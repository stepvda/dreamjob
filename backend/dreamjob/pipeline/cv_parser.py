"""CV parser for DOCX and PDF uploads (FR-103, NFR-402, NFR-205).

A CV is free-form where a LinkedIn export is generated, so this module reads
it as a sequence of *labelled lines* - text plus the three signals that a CV
actually carries: whether the line is a heading, whether it opens with an
emphasised lead, and whether it is a bullet.  DOCX gives those signals exactly
(bold runs, list styles, table layout); PDF approximates them from font size
and weight.  One section builder then consumes either.

The result lands in the FR-102 schema defined by ``linkedin_pdf`` so that
``profile_intake`` can diff the two sources field by field.

Two details of the product owner's own CV drove the design and are covered by
the unit tests: it is a single-table layout whose header cell holds the
contact block, and the photograph lives at ``word/media/*.jpg`` rather than in
any paragraph - the apply phase needs that file on disk (FR-103).
"""

from __future__ import annotations

import logging
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pdfplumber

from dreamjob.pipeline.linkedin_pdf import (
    MONTH_YEAR_PATTERN,
    ParsedDocument,
    empty_profile,
    extract_photo,
    parse_date_range,
    year_of,
)

log = logging.getLogger(__name__)

# CV headings vary far more than LinkedIn's; map generously onto FR-102 keys.
_CV_HEADERS: dict[str, str] = {
    "profile": "summary",
    "summary": "summary",
    "professional summary": "summary",
    "about": "summary",
    "about me": "summary",
    "objective": "summary",
    "personal statement": "summary",
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "employment history": "experience",
    "career history": "experience",
    "current training": "experience",
    "education": "education",
    "education training": "education",
    "academic background": "education",
    "skills": "top_skills",
    "skills capabilities": "top_skills",
    "core skills": "top_skills",
    "key skills": "top_skills",
    "technical skills": "top_skills",
    "core competencies": "top_skills",
    "areas of focus": "top_skills",
    "areas of expertise": "top_skills",
    "expertise": "top_skills",
    "languages": "languages",
    "language skills": "languages",
    "certifications": "certifications",
    "certificates": "certifications",
    "licenses certifications": "certifications",
    "training certifications": "certifications",
    "publications": "publications",
    "books": "publications",
    "papers": "publications",
    "projects": "projects",
    "selected projects": "projects",
    "side projects": "projects",
    "volunteering": "volunteering",
    "volunteer experience": "volunteering",
    "honors awards": "honors",
    "awards": "honors",
    "achievements": "honors",
    "interests": "interests",
    "hobbies": "interests",
    "references": "recommendations",
    "recommendations": "recommendations",
    "courses": "courses",
}

# A CV separates the fields of an entry with a padded middot; an unpadded one
# belongs to the field itself ("Business Unit Manager · Enterprise Content
# Management"), which is why the padded form is tried first.
_FIELD_SEP = re.compile(r"\s{2,}[·•|]\s{2,}|\s+[·•|]\s{2,}|\s{2,}[·•|]\s+")
_LOOSE_SEP = re.compile(r"\s*[·•|]\s*")
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(
    r"(?:\+|00)[\d][\d\s().-]{6,}"
    r"|\b0\d{1,2}[\s./-]?\d{2}[\s./-]?\d{2}[\s./-]?\d{2}\b"
)
_URL_RE = re.compile(r"(?:https?://)?(?:[\w-]+\.)+[a-z]{2,}(?:/[\w./#?=%&+-]*)?", re.IGNORECASE)

_TRAILING_RANGE_RE = re.compile(
    rf"((?:{MONTH_YEAR_PATTERN})\s*[-‐-―−]\s*(?:(?:{MONTH_YEAR_PATTERN})|[A-Za-zÀ-ſ]+)\.?)\s*$",
    re.IGNORECASE,
)

_ENTRY_LIST_KEYS = (
    "certifications", "publications", "projects", "honors",
    "volunteering", "recommendations", "courses",
)


@dataclass
class CvLine:
    """One line of a CV plus the signals that identify its role."""

    text: str
    is_heading: bool = False
    is_bullet: bool = False
    lead: str = ""          # emphasised opening run, when there is one
    tail: str = ""          # the rest of the line after that lead


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _docx_paragraph_line(para: Any) -> CvLine | None:
    text = para.text.replace("\xa0", " ").strip()
    if not text:
        return None
    style = (para.style.name if para.style is not None else "") or ""
    runs = [r for r in para.runs if r.text.strip()]
    bold_lead, rest = "", text
    if runs and runs[0].bold:
        bold_lead = runs[0].text.strip()
        rest = text[len(runs[0].text) :].strip() if text.startswith(runs[0].text.strip()) else text
        for run in runs[1:]:
            if not run.bold:
                break
            bold_lead = f"{bold_lead} {run.text.strip()}".strip()
    all_bold = bool(runs) and all(r.bold for r in runs)
    is_bullet = "list" in style.lower() or text[0] in "-–—•*"
    # A CV heading is a short, fully emphasised, upper-case line.
    is_heading = (
        not is_bullet
        and all_bold
        and len(text) <= 48
        and "\t" not in text
        and (text == text.upper() or style.lower().startswith("heading"))
    )
    return CvLine(
        text=re.sub(r"^[-–—•*]\s*", "", text),
        is_heading=is_heading,
        is_bullet=is_bullet,
        lead=bold_lead,
        tail=rest,
    )


def _docx_lines(path: Path) -> tuple[list[CvLine], list[str]]:
    """Body lines in reading order, plus the header/contact block."""
    import docx  # noqa: PLC0415 - optional at import time, required here

    document = docx.Document(str(path))
    header_block: list[str] = []
    lines: list[CvLine] = []

    # The contact block of a table-layout CV sits in the first table.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for chunk in cell.text.splitlines():
                    chunk = chunk.replace("\xa0", " ").strip()
                    if chunk:
                        header_block.append(chunk)
        if header_block:
            break

    for para in document.paragraphs:
        line = _docx_paragraph_line(para)
        if line is not None:
            lines.append(line)
    return lines, header_block


def extract_docx_photo(path: Path) -> bytes | None:
    """Largest embedded image in ``word/media`` - the portrait, in practice."""
    try:
        with zipfile.ZipFile(path) as zf:
            media = [
                n for n in zf.namelist()
                if n.startswith("word/media/")
                and n.lower().endswith((".jpg", ".jpeg", ".png", ".gif", ".bmp"))
            ]
            if not media:
                return None
            best = max(media, key=lambda n: zf.getinfo(n).file_size)
            return zf.read(best)
    except (zipfile.BadZipFile, KeyError, OSError):
        log.debug("No embedded image in %s", path, exc_info=True)
        return None


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_lines(path: Path) -> tuple[list[CvLine], list[str], str]:
    lines: list[CvLine] = []
    raw: list[str] = []
    records: list[tuple[float, str, bool]] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            raw.append(page.extract_text() or "")
            words = page.extract_words(extra_attrs=["size", "fontname"])
            buckets: dict[int, list[dict]] = {}
            for w in words:
                buckets.setdefault(int(round(w["top"] / 2.5)), []).append(w)
            for _, group in sorted(buckets.items()):
                group.sort(key=lambda w: w["x0"])
                text = " ".join(w["text"] for w in group).strip()
                # PDF text layers carry stray single glyphs (a rendered tab,
                # a bullet artefact); they carry no content.
                if len(text) < 2:
                    continue
                size = round(max(float(w["size"]) for w in group), 2)
                bold = any("bold" in str(w.get("fontname", "")).lower() for w in group)
                records.append((size, text, bold))

    body_size = Counter(s for s, _, _ in records).most_common(1)[0][0] if records else 10.0
    # A PDF CV has no table cell to hold the contact block: everything above
    # the first section heading is the header.
    first_heading = next(
        (i for i, (_, t, _) in enumerate(records) if _normalise(t) in _CV_HEADERS),
        min(len(records), 6),
    )
    header_block = [t for _, t, _ in records[:first_heading]]
    for size, text, bold in records[first_heading:]:
        stripped = re.sub(r"^[-–—•*]\s*", "", text)
        is_bullet = bool(re.match(r"^[-–—•*]\s+", text))
        is_heading = (
            not is_bullet
            and len(text) <= 48
            and (size > body_size + 0.5 or (bold and text == text.upper()))
            and _normalise(text) in _CV_HEADERS
            or (not is_bullet and len(text) <= 48 and size > body_size + 1.5)
        )
        lines.append(CvLine(text=stripped, is_heading=is_heading, is_bullet=is_bullet))
    return lines, header_block, "\n".join(raw)


# ---------------------------------------------------------------------------
# Section building
# ---------------------------------------------------------------------------


def _normalise(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", text.lower())).strip()


def _split_fields(text: str) -> list[str]:
    parts = [p.strip() for p in _FIELD_SEP.split(text) if p.strip()]
    if len(parts) < 2:
        parts = [p.strip() for p in _LOOSE_SEP.split(text) if p.strip()]
    return parts


def _split_dates(text: str) -> tuple[str, str | None]:
    """A CV puts the date range after a tab or at the end of the entry line."""
    if "\t" in text:
        head, _, tail = text.partition("\t")
        return head.strip(), tail.strip() or None
    m = _TRAILING_RANGE_RE.search(text)
    if m:
        return text[: m.start()].strip(" ,;·•|"), m.group(1).strip()
    return text.strip(), None


def _read_contact(header_block: list[str], doc: ParsedDocument) -> None:
    contact = doc.sections["contact"]
    for i, line in enumerate(header_block):
        email = _EMAIL_RE.search(line)
        if email and not contact["email"]:
            contact["email"] = email.group(0)
            doc.confidence("contact.email", 0.95)
        phone = _PHONE_RE.search(line)
        if phone and not contact["phone"]:
            contact["phone"] = phone.group(0).strip()
            doc.confidence("contact.phone", 0.9)
        for url in _URL_RE.findall(_EMAIL_RE.sub(" ", line)):
            url = url.rstrip(".,;")
            if "@" in url:
                continue
            if "linkedin.com/in/" in url.lower() and not contact["linkedin_url"]:
                contact["linkedin_url"] = url
                doc.confidence("contact.linkedin_url", 0.92)
            if url not in [w["url"] for w in contact["websites"]]:
                contact["websites"].append({"url": url, "label": None})
        if i == 0 and not contact["name"]:
            contact["name"] = line.strip()
            doc.confidence("contact.name", 0.9)
        elif i == 1 and not contact["headline"]:
            contact["headline"] = line.strip()
            doc.confidence("contact.headline", 0.85)
        elif not contact["location"]:
            first = re.split(r"\s*[|·•]\s*", line)[0].strip()
            if re.fullmatch(r"[A-Za-zÀ-ſ .'’-]{2,40},\s*[A-Za-zÀ-ſ .'’-]{2,40}", first):
                contact["location"] = first
                doc.confidence("contact.location", 0.8)


def _flush_prose(buffer: list[str], target: dict[str, Any], key: str) -> None:
    text = "\n".join(buffer).strip()
    if text:
        target[key] = ((target.get(key) or "") + "\n" + text).strip()


def _read_experience_entry(line: CvLine) -> dict[str, Any]:
    head, dates = _split_dates(line.text)
    parts = _split_fields(head)
    title = parts[0] if parts else head
    company = parts[1] if len(parts) > 1 else None
    location = parts[2] if len(parts) > 2 else None
    rng = parse_date_range(dates) if dates else None
    return {
        "company": company, "title": title, "start": rng.start if rng else None,
        "end": rng.end if rng else None, "current": bool(rng and rng.current),
        "location": location, "description": None,
        "duration_raw": dates, "company_duration": None,
    }


def _read_education_entry(line: CvLine) -> dict[str, Any]:
    head, dates = _split_dates(line.text)
    parts = _split_fields(head)
    rng = parse_date_range(dates) if dates else None
    return {
        "school": parts[1] if len(parts) > 1 else (parts[0] if parts else head),
        "degree": parts[0] if len(parts) > 1 else None,
        "field": None,
        "location": parts[2] if len(parts) > 2 else None,
        "start_year": year_of(rng.start) if rng else None,
        "end_year": year_of(rng.end) if rng else None,
        "description": None,
    }


def _build_sections(lines: list[CvLine], doc: ParsedDocument) -> None:
    section = "summary"
    recognised: list[str] = []
    entry: dict[str, Any] | None = None
    bullets: list[str] = []
    prose: list[str] = []
    unknown_title: str | None = None

    def close_entry() -> None:
        nonlocal entry, bullets
        if entry is not None:
            body = "\n".join(bullets).strip()
            if "school" in entry:
                entry["description"] = body or None
            elif "detail" in entry:
                entry["detail"] = ((entry.get("detail") or "") + "\n" + body).strip() or None
            else:
                entry["description"] = body or None
        entry, bullets = None, []

    def close_section() -> None:
        nonlocal prose, unknown_title
        close_entry()
        if prose:
            if section == "summary":
                _flush_prose(prose, doc.sections, "summary")
            elif section == "other" and unknown_title:
                doc.sections["other"][unknown_title] = "\n".join(prose).strip()
            prose = []
        unknown_title = None

    for line in lines:
        if line.is_heading:
            close_section()
            key = _CV_HEADERS.get(_normalise(line.text))
            if key is None:
                section, unknown_title = "other", line.text.strip()
            else:
                section = key
                recognised.append(key)
            continue

        if section == "summary":
            prose.append(line.text)
            doc.confidence("summary", 0.9)
            continue

        if section == "experience":
            if line.is_bullet:
                bullets.append(line.text)
            else:
                close_entry()
                entry = _read_experience_entry(line)
                doc.sections["experience"].append(entry)
            continue

        if section == "education":
            # A bullet, or a sentence, continues the study above it; anything
            # else starts a new one.
            if entry is not None and (line.is_bullet or line.text.endswith((".", ";"))):
                bullets.append(line.text)
            else:
                close_entry()
                entry = _read_education_entry(line)
                doc.sections["education"].append(entry)
            continue

        if section == "top_skills":
            _read_skill_line(line, doc)
            continue

        if section == "languages":
            _read_language_line(line, doc)
            continue

        if section == "interests":
            doc.sections["interests"].append(line.text)
            continue

        if section in _ENTRY_LIST_KEYS:
            if line.is_bullet and entry is not None:
                bullets.append(line.text)
            else:
                close_entry()
                head, dates = _split_dates(line.text)
                parts = _split_fields(head)
                rng = parse_date_range(dates) if dates else None
                entry = {
                    "title": parts[0] if parts else head,
                    "detail": " · ".join(parts[1:]) or None,
                    "start": rng.start if rng else None,
                    "end": rng.end if rng else None,
                }
                doc.sections[section].append(entry)
            continue

        prose.append(line.text)

    close_section()
    if not recognised:
        # No heading was recognised at all: hand the whole document to the LLM.
        doc.unsegmented["whole document"] = "\n".join(ln.text for ln in lines)

    for i, _ in enumerate(doc.sections["experience"]):
        doc.confidence(f"experience.{i}.company", 0.85)
        doc.confidence(f"experience.{i}.title", 0.88)
        doc.confidence(f"experience.{i}.dates", 0.8)
    for i, _ in enumerate(doc.sections["education"]):
        doc.confidence(f"education.{i}.school", 0.85)


def _read_skill_line(line: CvLine, doc: ParsedDocument) -> None:
    """A CV skills line is either one skill, or a category and its members."""
    text = line.text
    categorised = bool(line.lead and line.tail)
    body = line.tail if categorised else text
    # A bullet written as a sentence is an "area of focus", not a skill label;
    # splitting it on commas would manufacture labels nothing can normalise.
    if not categorised and len(body.split()) >= 7 and body.rstrip().endswith("."):
        doc.sections["other"].setdefault("areas_of_focus", [])
        doc.sections["other"]["areas_of_focus"].append(text)
        return
    for chunk in _split_outside_parens(body):
        label = chunk.strip(" .;")
        if 1 < len(label) <= 80:
            doc.sections["top_skills"].append(label)
    idx = len(doc.sections["top_skills"]) - 1
    if idx >= 0:
        doc.confidence(f"top_skills.{idx}", 0.8)


def _split_outside_parens(text: str) -> list[str]:
    """Split a skills line on separators, keeping bracketed lists intact."""
    parts: list[str] = []
    buf: list[str] = []
    depth = 0
    for ch in text:
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth = max(0, depth - 1)
        if depth == 0 and ch in ",;•":
            parts.append("".join(buf))
            buf = []
            continue
        buf.append(ch)
    parts.append("".join(buf))
    out: list[str] = []
    for part in parts:
        out.extend(p for p in re.split(r"\s{2,}", part) if p.strip())
    return out


def _read_language_line(line: CvLine, doc: ParsedDocument) -> None:
    text = line.text
    pairs = re.findall(
        r"([A-Za-zÀ-ſ]{3,20})\s*[:(]?\s*"
        r"(native|bilingual|fluent|full professional|professional|working|"
        r"elementary|basic|beginner|intermediate|advanced|mother tongue|"
        r"moedertaal|courant|maternelle)\)?",
        text,
        re.IGNORECASE,
    )
    if pairs:
        for language, proficiency in pairs:
            doc.sections["languages"].append(
                {"language": language.strip().title(), "proficiency": proficiency.strip().lower()}
            )
        doc.confidence(f"languages.{len(doc.sections['languages']) - 1}.language", 0.85)
        return
    for chunk in re.split(r"[,;/]|\s{2,}", text):
        label = chunk.strip(" .;")
        if 2 < len(label) <= 30:
            doc.sections["languages"].append({"language": label, "proficiency": None})


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def parse_cv(cv_path: str | Path) -> ParsedDocument:
    """Parse a DOCX or PDF CV into the FR-102 profile schema (FR-103)."""
    path = Path(cv_path)
    doc = ParsedDocument(source="cv", sections=empty_profile())
    suffix = path.suffix.lower()
    if suffix == ".docx":
        lines, header_block = _docx_lines(path)
        doc.raw_text = "\n".join(header_block + [ln.text for ln in lines])
        doc.photo = extract_docx_photo(path)
    elif suffix == ".pdf":
        lines, header_block, raw = _pdf_lines(path)
        doc.raw_text = raw
        doc.photo = extract_photo(path)
    elif suffix in {".doc", ".rtf", ".odt"}:
        raise ValueError(
            f"{suffix} is not supported; export the CV as .docx or .pdf and upload it again."
        )
    else:
        raise ValueError(f"Unsupported CV format {suffix!r}: upload a .pdf or .docx file.")

    _read_contact(header_block, doc)
    _build_sections(lines, doc)
    if not doc.sections["experience"]:
        doc.warnings.append("No experience entries were recognised in the CV.")
        doc.unsegmented["whole document"] = doc.raw_text
    return doc


def save_photo(photo: bytes, seeker_id: str, uploads_dir: Path) -> str:
    """Write the portrait to ``data/uploads/<seeker>/photo.jpg`` (FR-103).

    The apply phase reads it back from ``profile_version.photo_path``, so the
    file name is stable and the extension follows the actual bytes.
    """
    ext = ".jpg"
    if photo[:8] == b"\x89PNG\r\n\x1a\n":
        ext = ".png"
    elif photo[:6] in (b"GIF87a", b"GIF89a"):
        ext = ".gif"
    target_dir = uploads_dir / seeker_id
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"photo{ext}"
    target.write_bytes(photo)
    return str(target)
