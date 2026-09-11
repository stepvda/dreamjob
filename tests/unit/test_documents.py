"""Generated application documents (FR-321..331, NFR-206, NFR-302, CR-405, RK-03).

Everything here runs against a throwaway SQLite file with no network and no
LLM: the environment fixture clears the provider key and records no
``llm_transfer`` consent, so every generator takes its deterministic path and
the tests assert on what the system produces when the model is unavailable -
which is also what CR-405 says it must still be safe to send.

The consistency tests work the other way round: they take a correctly generated
document and poison it the way a hallucinating model would, then assert the
validator refuses it.
"""

from __future__ import annotations

import base64
import os
import secrets
from collections.abc import Iterator
from pathlib import Path

import pytest
from dreamjob.config import get_settings
from fastapi import FastAPI
from fastapi.testclient import TestClient

_ENV_KEYS = (
    "DREAMJOB_DATA_DIR", "DREAMJOB_DB_PATH", "DREAMJOB_MASTER_KEY", "DREAMJOB_ENV",
    "DEEPSEEK_API_KEY", "DREAMJOB_LOCAL_LLM_BASE_URL",
)


@pytest.fixture(autouse=True)
def isolated_db(tmp_path: Path) -> Iterator[None]:
    saved = {k: os.environ.get(k) for k in _ENV_KEYS}
    os.environ["DREAMJOB_DATA_DIR"] = str(tmp_path / "data")
    os.environ["DREAMJOB_DB_PATH"] = str(tmp_path / "data" / "test.db")
    os.environ["DREAMJOB_MASTER_KEY"] = base64.urlsafe_b64encode(secrets.token_bytes(32)).decode()
    os.environ["DREAMJOB_ENV"] = "development"
    # No provider, so every generator must take its non-LLM path (NFR-104).
    os.environ["DEEPSEEK_API_KEY"] = ""
    os.environ["DREAMJOB_LOCAL_LLM_BASE_URL"] = ""
    get_settings.cache_clear()

    from dreamjob.db.migrator import migrate

    migrate()
    yield

    for key, value in saved.items():
        if value is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = value
    get_settings.cache_clear()


# ---------------------------------------------------------------------------
# Fixtures: one job seeker, one company, one vacancy, one speculative opening
# ---------------------------------------------------------------------------

SECTIONS = {
    "contact": {
        "name": "Stephane van der Aa",
        "headline": "Data & AI",
        "email": "stephane@stepvda.com",
        "phone": "+32 470 11 22 33",
        "location": "Gent",
        "linkedin_url": "linkedin.com/in/svda",
        "websites": [],
    },
    "summary": "Twintig jaar ervaring in data en analytics.",
    "top_skills": ["Python", "SQL"],
    "languages": [{"language": "Nederlands", "level": "moedertaal"}],
    "experience": [
        {
            "company": "Acme N.V.", "title": "Head of Data", "start": "2019-05", "end": None,
            "current": True, "location": "Gent",
            "description": "Team van 12 opgebouwd.\nKosten met 30% verlaagd.",
        },
        {
            "company": "Beta BV", "title": "Lead Data Engineer", "start": "2014-01",
            "end": "2019-04", "location": "Antwerpen",
            "description": "Realtime pijplijn gebouwd.",
        },
    ],
    "education": [
        {"school": "Universiteit Gent", "degree": "MSc", "field": "Informatica",
         "start_year": "2003", "end_year": "2008", "description": None}
    ],
    "certifications": [{"title": "Azure Data Engineer", "start": "2022"}],
}


def _photo(tmp_path: Path) -> str:
    """A real JPEG, so the templates exercise the FR-322 photograph path."""
    from PIL import Image

    path = tmp_path / "photo.jpg"
    Image.new("RGB", (300, 300), (40, 60, 90)).save(path, "JPEG")
    return str(path)


def seed(photo_path: str | None = None) -> dict[str, str]:
    from dreamjob.db.connection import insert_row, to_json, utcnow

    now = utcnow()
    seeker = insert_row(
        "job_seeker",
        {"email": "stephane@stepvda.com", "display_name": "Stephane van der Aa",
         "locale": "nl", "created_at": now, "updated_at": now},
    )
    version = insert_row(
        "profile_version",
        {"job_seeker_id": seeker, "version": 7, "sections": to_json(SECTIONS),
         "photo_path": photo_path, "source_note": "merged", "created_at": now},
    )
    for skill in ("Python", "SQL", "Databricks", "Teamleiding"):
        insert_row(
            "profile_skill",
            {"job_seeker_id": seeker, "profile_version_id": version, "raw_label": skill,
             "normalised_label": skill, "proficiency": 4},
        )
    directives = insert_row(
        "directive_set", {"job_seeker_id": seeker, "name": "Standaard", "created_at": now}
    )
    campaign = insert_row(
        "campaign",
        {"job_seeker_id": seeker, "directive_set_id": directives, "profile_version_id": version,
         "name": "Najaar", "created_at": now},
    )
    company = insert_row(
        "company",
        {"normalised_name": "zenith analytics", "name": "Zenith Analytics NV", "country": "BE",
         "domain": "zenith.example",
         "business_summary": "Zenith bouwt analytics software voor de zorgsector.",
         "size_fte": 180, "size_band": "100-250", "stage": "scaleup", "trajectory": "growing",
         "values_culture": to_json([{"name": "Ownership"}]),
         "collected_at": now, "refreshed_at": now, "source": "website"},
    )
    for index, (revenue, ebitda, headcount) in enumerate(
        [(8e6, 6e5, 80), (9.5e6, 8e5, 95), (11e6, 9e5, 120), (13.2e6, 1.1e6, 150),
         (15.8e6, 1.4e6, 180)]
    ):
        insert_row(
            "financial_year",
            {"company_id": company, "fiscal_year": 2020 + index, "currency": "EUR",
             "revenue": revenue, "ebitda": ebitda, "headcount_fte": headcount,
             "source": "NBB", "collected_at": now},
        )
    insert_row(
        "financial_analysis",
        {"company_id": company, "years_covered": to_json([2020, 2024]), "revenue_cagr": 0.18,
         "trajectory": "growing", "ability_to_pay": 78, "computed_at": now},
    )
    insert_row(
        "hiring_signal",
        {"company_id": company, "signal_type": "funding", "description": "Serie B",
         "occurred_at": "2026-07-02", "collected_at": now},
    )
    contact = insert_row(
        "contact",
        {"company_id": company, "full_name": "Els Peeters", "role_title": "CTO",
         "email": "els.peeters@zenith.example", "email_validation": "valid",
         "collected_at": now, "confidence": 0.9},
    )
    vacancy = insert_row(
        "vacancy",
        {"company_id": company, "title": "Head of Data",
         "description": "Wij zoeken een Head of Data.",
         "required_skills": to_json(["Python", "Databricks", "Teamleiding"]),
         "desirable_skills": to_json(["dbt"]), "location": "Gent", "language": "nl",
         "collected_at": now},
    )
    opportunity = insert_row(
        "opportunity",
        {"job_seeker_id": seeker, "campaign_id": campaign, "company_id": company,
         "vacancy_id": vacancy, "kind": "vacancy", "title": "Head of Data", "language": "nl",
         "selected": 1, "description": "Wij zoeken een Head of Data.",
         "required_skills": to_json(["Python", "Databricks", "Teamleiding"]),
         "created_at": now, "updated_at": now},
    )
    speculative = insert_row(
        "opportunity",
        {"job_seeker_id": seeker, "campaign_id": campaign, "company_id": company,
         "kind": "speculative", "title": "Director of Data Platform", "language": "nl",
         "selected": 1, "speculative_rationale": "Zenith groeit snel.", "plausibility": 0.62,
         "created_at": now, "updated_at": now},
    )
    return {
        "seeker": seeker, "campaign": campaign, "company": company, "contact": contact,
        "opportunity": opportunity, "speculative": speculative, "profile_version": version,
    }


def _inputs(ids: dict[str, str], key: str = "opportunity") -> dict:
    from dreamjob.db.repositories import applications as repo

    return repo.generation_inputs(ids["seeker"], ids[key])


# ---------------------------------------------------------------------------
# FR-322: templates, formats, languages, photograph
# ---------------------------------------------------------------------------


def test_template_registry_discovers_both_templates() -> None:
    from dreamjob.documents.templates import available, get_template

    keys = [t.key for t in available()]
    assert {"classic", "modern"} <= set(keys)
    assert keys[0] == "classic", "the default template must sort first"
    assert get_template("does-not-exist").TEMPLATE_KEY == "classic"


def test_cv_generation_without_llm(tmp_path: Path) -> None:
    """FR-322: DOCX and PDF, from the profile alone, with the photograph."""
    from docx import Document
    from dreamjob.documents.cv_generator import generate_cv

    ids = seed(photo_path=_photo(tmp_path))
    inputs = _inputs(ids)

    for template in ("classic", "modern"):
        result = generate_cv(
            inputs, output_dir=tmp_path / template, language="nl", template=template
        )
        assert result.template == template
        assert not result.tailored_by_llm
        docx_path, pdf_path = Path(result.docx_path), Path(result.pdf_path)
        assert docx_path.is_file() and pdf_path.is_file()
        assert pdf_path.stat().st_size > 3000

        document = Document(str(docx_path))
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        assert "Stephane van der Aa" in text
        assert "Acme N.V." in text
        # The CV must carry the photograph the profile holds (FR-322).
        assert any("image" in part.content_type for part in document.part.package.iter_parts())


def test_the_contact_line_prints_the_address_not_the_object() -> None:
    """FR-322: a website the profile holds reaches the CV as its address.

    ``linkedin_pdf``, ``cv_parser`` and the profile merge all write
    ``contact.websites`` as ``{"url": ..., "label": ...}``.  The CV header used
    to take ``str()`` of one of those, which printed a Python dict repr into
    the contact line of the document that gets attached and sent.
    """
    from dreamjob.db.connection import to_json, update_row
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed()
    sections = {
        **SECTIONS,
        "contact": {
            **SECTIONS["contact"],
            "websites": [
                {"url": "www.linkedin.com/in/svda", "label": "LinkedIn"},
                {"url": "stepvda.example.com", "label": None},
                "plain.example.com",
            ],
        },
    }
    update_row("profile_version", ids["profile_version"], {"sections": to_json(sections)})

    document = build_base_document(_inputs(ids), language="nl")
    assert document.contact.websites == [
        "www.linkedin.com/in/svda",
        "stepvda.example.com",
        "plain.example.com",
    ]
    assert not any("{" in site for site in document.contact.websites)
    assert "label" not in " ".join(document.contact.websites)

    # ``linkedin_url`` is copied out of ``websites`` upstream, so the header
    # line must not print the same profile address twice.
    lines = document.contact.lines()
    assert lines.count("linkedin.com/in/svda") + lines.count("www.linkedin.com/in/svda") == 1
    assert len(lines) == len(set(line.casefold() for line in lines))


def test_a_model_that_answers_with_a_list_still_produces_a_cv(tmp_path: Path) -> None:
    """NFR-104: a badly shaped answer costs the tailoring, not the package.

    ``generate_cv`` degrades on ``LLMError``/``ValueError``/``KeyError``/
    ``TypeError``.  A model that returns a bare JSON list used to raise
    ``AttributeError`` inside ``apply_tailoring`` instead, which is not in that
    tuple, so the whole package generation failed on one bad answer.
    """
    import pytest as _pytest
    from dreamjob.documents.cv_generator import apply_tailoring, build_base_document, generate_cv

    ids = seed()
    inputs = _inputs(ids)

    with _pytest.raises(ValueError, match="not an object"):
        apply_tailoring(build_base_document(inputs, language="nl"), ["headline", "summary"])

    class ListAnsweringClient:
        def complete_json(self, *args: object, **kwargs: object) -> object:
            return ["headline", "summary"]

    result = generate_cv(inputs, output_dir=tmp_path, language="nl", llm=ListAnsweringClient())
    assert Path(result.pdf_path).is_file()
    assert not result.tailored_by_llm
    assert any("Tailoring unavailable" in note for note in result.notes)


def test_cv_language_follows_the_opportunity() -> None:
    """FR-322 / NFR-501: nl, fr, en and de all produce localised headings."""
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed()
    inputs = _inputs(ids)
    expected = {"nl": "Werkervaring", "fr": "Expérience professionnelle",
                "en": "Experience", "de": "Berufserfahrung"}
    for language, heading in expected.items():
        document = build_base_document(inputs, language=language)
        assert document.section_label("experience") == heading
        assert document.experience[0].period(language)


def test_do_not_disclose_is_absolute(tmp_path: Path) -> None:
    """FR-106: a flagged field never reaches the document, photograph included."""
    from dreamjob.db.connection import insert_row, utcnow
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed(photo_path=_photo(tmp_path))
    for field_path in ("contact.phone", "photo", "experience.1"):
        insert_row(
            "disclosure_flag",
            {"job_seeker_id": ids["seeker"], "field_path": field_path,
             "do_not_disclose": 1, "created_at": utcnow()},
        )
    document = build_base_document(_inputs(ids), language="nl")

    assert document.contact.phone is None
    assert document.photo_path is None
    assert [job.company for job in document.experience] == ["Acme N.V."]
    assert "+32 470" not in document.plain_text()


def test_do_not_disclose_never_reaches_the_model_payload() -> None:
    """FR-106 is absolute: a flagged field must not leave for the model either."""
    from dreamjob.documents.consistency import profile_facts
    from dreamjob.documents.cv_generator import Disclosure

    ids = seed()
    inputs = _inputs(ids)
    inputs["do_not_disclose"] = {"contact.phone"}
    version = dict(inputs["profile_version"])
    sections = dict(version["sections"])
    sections["contact"] = {**(sections.get("contact") or {}), "phone": "+32 470 11 22 33"}
    version["sections"] = sections
    inputs["profile_version"] = version

    facts = profile_facts(inputs)
    assert "470 11 22 33" not in facts.corpus

    redacted = Disclosure({"contact.phone"}).redact(sections)
    assert "phone" not in redacted["contact"]


# ---------------------------------------------------------------------------
# FR-322 / RK-03: the factual-consistency validator
# ---------------------------------------------------------------------------


def _report(document, ids, key: str = "opportunity"):
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import consistency

    return consistency.run_checks(
        document,
        _inputs(ids, key),
        repo.provenance_corpus(ids["seeker"], ids[key]),
        use_judge=False,
    )


def test_untailored_cv_passes_its_own_check() -> None:
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed()
    report = _report(build_base_document(_inputs(ids), language="nl"), ids)
    assert report.status == "pass", [f.as_dict() for f in report.findings]
    assert report.leak_status == "pass", [f.as_dict() for f in report.leaks]
    assert report.claims_checked > 10


@pytest.mark.parametrize(
    ("mutate", "kind"),
    [
        (lambda d: setattr(d.experience[0], "company", "Globex International"), "employer"),
        (lambda d: setattr(d.experience[0], "title", "Chief Executive Officer"), "title"),
        (lambda d: setattr(d.experience[1], "start", "2011-03"), "period"),
        (lambda d: setattr(d.education[0], "school", "Harvard Business School"), "education"),
        (lambda d: d.skills.append("Kubernetes"), "skill"),
        (lambda d: setattr(d, "summary", "Ik leidde een team van 45 mensen."), "figure"),
    ],
)
def test_consistency_rejects_unsupported_claims(mutate, kind: str) -> None:
    """RK-03: each of these is a hallucination the deterministic checks must catch."""
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed()
    document = build_base_document(_inputs(ids), language="nl")
    mutate(document)
    report = _report(document, ids)

    assert report.status == "fail"
    assert kind in {f.kind for f in report.findings}
    assert all(f.signal == "deterministic" for f in report.findings)


def test_normalised_skills_count_as_profile_facts() -> None:
    """FR-107 skills live in their own table; prose citing one is not a claim."""
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed()
    document = build_base_document(_inputs(ids), language="nl")
    document.summary = "Werkt dagelijks met Databricks en Python."
    report = _report(document, ids)

    assert report.status == "pass", [f.as_dict() for f in report.findings]
    assert "Databricks" not in {f.claim for f in report.findings}


def test_judge_findings_only_fail_on_unsupported_figures_or_names() -> None:
    """RK-03: the judge is a signal.  Only a mechanical miss fails a document."""
    from dreamjob.documents.consistency import profile_facts, unsupported_tokens

    facts = profile_facts(_inputs(seed()))

    # A paraphrase of profile material has nothing the profile lacks.
    assert unsupported_tokens("Geeft leiding aan een team van 12 bij Acme N.V.", facts) == []
    # A fabricated employer and a fabricated figure do.
    assert "Globex International" in unsupported_tokens("Werkte bij Globex International.", facts)
    assert "45" in unsupported_tokens("Leidde 45 mensen.", facts)


def test_the_judge_cannot_fail_a_package_on_its_own() -> None:
    """FR-322/RK-03: a bare judge ``high`` with nothing mechanically missing is capped.

    The severity cap used ``min`` over a most-severe-first map, so ``high``
    stayed high and a model that disliked the tone could block a dispatch on
    its own - the invariant this module states.
    """
    from dreamjob.documents.consistency import judge, profile_facts

    facts = profile_facts(_inputs(seed()))

    class Dislikes:
        def complete_json(self, *args: object, **kwargs: object) -> object:
            return {
                "unsupported": [
                    {"quote": "a great fit for the team", "severity": "high", "why": "tone"}
                ]
            }

    findings, error = judge("CV text", facts, Dislikes())

    assert error is None
    assert findings, "the finding is still reported"
    assert findings[0].severity == "medium", "but it cannot fail the document"


def test_the_judge_tolerates_a_bare_json_array() -> None:
    """complete_json returns whatever parsed; a list must degrade, not raise."""
    from dreamjob.documents.consistency import judge, profile_facts

    facts = profile_facts(_inputs(seed()))

    class AnswersAList:
        def complete_json(self, *args: object, **kwargs: object) -> object:
            return ["not", "an", "object"]

    findings, error = judge("CV text", facts, AnswersAList())

    assert findings == []
    assert error is not None


def test_judge_allows_the_company_and_role_the_application_is_for() -> None:
    """FR-322: the judge and check_free_text must agree about one name.

    A tailored CV names its target, and ``check_free_text`` is given that name
    in ``allow``.  Escalating the judge's quote on the same name would fail the
    package on a word the sibling check just permitted - and would let the
    judge fail a document on its own, which this module forbids.
    """
    from dreamjob.documents.consistency import profile_facts, unsupported_tokens

    facts = profile_facts(_inputs(seed()))
    quote = "This background fits the Senior Fullstack Engineer role at Solactive."
    allow = {"Solactive", "Senior Fullstack Engineer"}

    assert unsupported_tokens(quote, facts) != []  # without the set, it escalates
    assert unsupported_tokens(quote, facts, allow=allow) == []
    # The allowance is narrow: another employer is still caught.
    assert "Globex International" in unsupported_tokens(
        "Werkte bij Globex International.", facts, allow=allow
    )


def test_judge_does_not_read_a_skill_headline_as_an_invented_name() -> None:
    """E2E_1500 section 8.8: 46 of 70 packages failed on CV headline prose.

    The judge escalated maximal runs of capitalised words - "Enterprise
    Architecture Expertise", "Data-Driven Product Experience" - because words
    like *Expertise* are not in the profile.  A real employer name still fails.
    """
    from dreamjob.documents.consistency import profile_facts, unsupported_tokens

    facts = profile_facts(_inputs(seed()))
    for headline in (
        "Enterprise Architecture Expertise",
        "Data-Driven Product Experience",
        "Software Engineer & Technical Leader",
        "German and English",
        "Jahren Erfahrung",
    ):
        assert unsupported_tokens(headline, facts) == [], headline
    assert "Globex International" in unsupported_tokens(
        "Werkte bij Globex International.", facts
    )


def test_judge_escalation_honours_german_noun_capitalisation() -> None:
    """FR-322: in German, a capitalised word is a noun, not a name.

    ``capitalised_phrases`` already knows this - "German capitalises every noun,
    so there a run has to be more than one word ... otherwise the scan would
    report most of the document" - and ``check_free_text`` passes the document's
    language so it applies.  ``unsupported_tokens`` did not, so the judge
    escalation read *Erfahrung*, *Systeme* and *Entwickler* as invented
    employers and failed German CVs by construction.
    """
    from dreamjob.documents.consistency import profile_facts, unsupported_tokens

    facts = profile_facts(_inputs(seed()))
    german = "Ein Jahrzehnt Erfahrung mit verteilten Systemen und Kern-APIs für Entwickler."

    flagged_as_english = unsupported_tokens(german, facts)
    flagged_as_german = unsupported_tokens(german, facts, language="de")
    assert len(flagged_as_german) < len(flagged_as_english)
    for ordinary_noun in ("Systemen", "Entwickler", "Kern-APIs"):
        assert ordinary_noun not in flagged_as_german

    # A real employer is still caught, in German too.
    assert "Globex International GmbH" in unsupported_tokens(
        "Arbeitete bei Globex International GmbH.", facts, language="de"
    )


def test_leak_scan_flags_another_persons_details() -> None:
    """NFR-206: content that is not this job seeker's must be reported."""
    from dreamjob.documents.cv_generator import build_base_document

    ids = seed()
    document = build_base_document(_inputs(ids), language="nl")
    document.contact.email = "someone.else@othercompany.com"
    report = _report(document, ids)

    assert report.leak_status == "fail"
    kinds = {f.kind for f in report.leaks}
    assert "leak_identity" in kinds
    assert "leak_email" in kinds


def test_leak_scan_allows_the_company_only_where_it_belongs() -> None:
    """The email may name the employer it is addressed to; the CV may not."""
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import consistency

    ids = seed()
    provenance = consistency.build_provenance(
        repo.provenance_corpus(ids["seeker"], ids["opportunity"])
    )
    text = "Ik schrijf u over Zenith Analytics NV."

    assert consistency.scan_leakage({"email": text}, provenance, with_context={"email"}) == []
    assert consistency.scan_leakage({"cv": text}, provenance, with_context=set())


def test_leak_scan_does_not_read_a_span_of_years_as_a_phone_number() -> None:
    """NFR-206: a date is not a leaked telephone number.

    ``2015-2017`` reduces to eight digits, which the phone pattern accepted -
    and a ``leak_phone`` finding is high severity, so ``leak_scan_status``
    failed and FR-324 refused approval with *no* override.  Every CV and
    motivation document dates its experience this way.
    """
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import consistency

    ids = seed()
    provenance = consistency.build_provenance(
        repo.provenance_corpus(ids["seeker"], ids["opportunity"])
    )
    for dated in ("Lead engineer 2015-2017 at Acme.", "Studied there (2015–2017).", "1999-2024"):
        leaks = consistency.scan_leakage({"cv": dated}, provenance, with_context=set())
        assert not [f for f in leaks if f.kind == "leak_phone"], dated

    # A real number is still reported.
    leaks = consistency.scan_leakage({"cv": "Reach me on +32 471 12 34 56."}, provenance,
                                     with_context=set())
    assert [f for f in leaks if f.kind == "leak_phone"]


# ---------------------------------------------------------------------------
# FR-329 / FR-330: the two job-seeker documents
# ---------------------------------------------------------------------------


def test_briefing_without_llm(tmp_path: Path) -> None:
    """FR-329: every named section, the five-year figures, and charts."""
    from dreamjob.documents.briefing import generate_briefing
    from pypdf import PdfReader

    ids = seed()
    result = generate_briefing(_inputs(ids), output_dir=tmp_path, language="nl")
    assert not result.used_llm

    reader = PdfReader(result.path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for heading in (
        "Bedrijfsprofiel", "Financiële analyse over vijf jaar", "Aanwervingssignalen",
        "Concurrenten", "De opening", "Waarschijnlijke gespreksonderwerpen", "Vragen om zelf",
    ):
        assert heading in text, f"{heading!r} missing from the briefing"
    assert "2024" in text and "15.80M EUR" in text
    # FR-329 asks for charts; reportlab draws them as vector content, so the
    # page carries drawing operators rather than an embedded image.
    assert "Omzet en resultaat" in text
    # FR-331: date and data versions on every page.
    assert "Profielversie 7" in text
    assert "Alleen voor de werkzoekende" in text


def test_briefing_is_regenerable_on_demand(tmp_path: Path) -> None:
    """FR-329: refreshable just before an interview, over the same path."""
    from dreamjob.documents.briefing import generate_briefing

    ids = seed()
    inputs = _inputs(ids)
    first = generate_briefing(inputs, output_dir=tmp_path, language="nl")
    second = generate_briefing(inputs, output_dir=tmp_path, language="nl")
    assert first.path == second.path
    assert Path(second.path).is_file()


def test_speculative_requirements_come_from_the_rationale(tmp_path: Path) -> None:
    """FR-330: a speculative opening still gets a requirement-by-requirement map."""
    from dreamjob.documents.motivation import generate_motivation, requirements

    ids = seed()
    inputs = _inputs(ids, "speculative")
    assert requirements(inputs) == ["Zenith groeit snel."]

    result = generate_motivation(inputs, output_dir=tmp_path, language="nl")
    assert [r["requirement"] for r in result.content["why_fit_job"]] == ["Zenith groeit snel."]
    assert Path(result.path).is_file()


def test_motivation_fallback_without_llm(tmp_path: Path) -> None:
    """FR-330: every requirement gets a row, and the ones with no evidence say so."""
    from dreamjob.documents.motivation import generate_motivation, requirements

    ids = seed()
    inputs = _inputs(ids)
    result = generate_motivation(inputs, output_dir=tmp_path, language="nl")

    assert not result.used_llm
    assert Path(result.path).is_file()
    rows = result.content["why_fit_job"]
    assert {r["requirement"] for r in rows} == set(requirements(inputs))
    strengths = {r["requirement"]: r["strength"] for r in rows}
    assert strengths["Python"] == "strong"
    assert strengths["dbt"] == "gap"
    assert "gap" in {r["strength"] for r in rows}
    assert result.as_dict()["gaps"] == ["dbt"]


# ---------------------------------------------------------------------------
# FR-321 / FR-323 / NFR-302: the introduction email
# ---------------------------------------------------------------------------


def test_speculative_email_asserts_no_vacancy() -> None:
    """FR-323: a spontaneous application proposes a role, it does not answer an ad."""
    from dreamjob.documents.cv_generator import build_base_document
    from dreamjob.documents.intro_email import compose_email, vacancy_assertions

    ids = seed()
    inputs = _inputs(ids, "speculative")
    document = build_base_document(inputs, language="nl")
    draft = compose_email(inputs, document, language="nl")

    assert draft.speculative
    assert draft.assertions == []
    assert "Spontane sollicitatie" in draft.subject
    assert "Director of Data Platform" in draft.body
    assert vacancy_assertions("Naar aanleiding van uw vacature schrijf ik u.", "nl")
    assert vacancy_assertions("I am applying for the advertised position.", "en")


def test_every_email_carries_the_objection_sentence() -> None:
    """NFR-302: legitimate interest requires a route to object, in every email."""
    from dreamjob.documents.cv_generator import build_base_document
    from dreamjob.documents.intro_email import OBJECTION_SENTENCE, compose_email

    ids = seed()
    for key in ("opportunity", "speculative"):
        inputs = _inputs(ids, key)
        for language in ("nl", "fr", "en", "de"):
            document = build_base_document(inputs, language=language)
            body = compose_email(inputs, document, language=language).body
            assert OBJECTION_SENTENCE[language] in body
            assert body.count(OBJECTION_SENTENCE[language]) == 1


def test_invented_figures_in_the_email_block_dispatch() -> None:
    """RK-03 names the CV *and* the email, and the email is what gets sent.

    A figure that is in neither the profile nor the opening is a hallucination
    wherever it appears, so the deterministic check runs over the email body as
    well and an edit that introduces one re-opens the gate (CR-405).
    """
    from dreamjob.documents import package as package_module

    ids = seed()
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    assert package["consistency_status"] == "pass"

    edited = package_module.edit(
        ids["seeker"],
        package["id"],
        {
            "email_body": "Beste Els Peeters,\n\nBij Acme N.V. verhoogde ik de omzet met "
            "480% en dat sinds 1997.\n\nMet vriendelijke groet,\nStephane van der Aa\n",
        },
    )
    assert edited["consistency_status"] == "fail"
    invented = {
        finding["claim"]
        for finding in (edited["consistency_report"] or {})["findings"]
        if finding["locus"] == "email" and finding["severity"] == "high"
    }
    assert {"480", "1997"} <= invented
    assert package_module.approve(ids["seeker"], [edited["id"]], actor="test")["approved"] == []


@pytest.mark.parametrize("language", ["en", "nl", "fr", "de"])
@pytest.mark.parametrize("opening", ["opportunity", "speculative"])
def test_generated_material_passes_its_own_checks(language: str, opening: str) -> None:
    """A scan that cries wolf over the system's own wording is worse than none.

    The greeting, the attachment line, the sign-off, the NFR-302 objection
    sentence and the month names in a date range are written by this package,
    not by a model, so none of them may be reported as content of unknown
    origin - in any of the four content languages (NFR-206, NFR-501).
    """
    from dreamjob.documents import package as package_module

    package = package_module.generate(
        (ids := seed())["seeker"],
        ids[opening],
        package_module.GenerationOptions(language=language),
        actor="test",
    )
    report = package["consistency_report"] or {}
    assert package["consistency_status"] == "pass", report.get("findings")
    assert package["leak_scan_status"] == "pass", report.get("leaks")
    assert package_module.approval_blockers(package) == []


# ---------------------------------------------------------------------------
# FR-321 / FR-324: the package
# ---------------------------------------------------------------------------


def test_package_generates_four_artefacts_and_attaches_only_the_cv(tmp_path: Path) -> None:
    from dreamjob.documents import package as package_module

    ids = seed(photo_path=_photo(tmp_path))
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")

    for key in ("cv_docx_path", "cv_pdf_path", "briefing_pdf_path", "motivation_pdf_path"):
        assert package[key] and Path(package[key]).is_file(), key
    assert package["email_subject"] and package["email_body"]
    assert package["consistency_status"] == "pass"
    assert package["leak_scan_status"] == "pass"
    # FR-331: the versions the documents were generated from.
    assert package["profile_version_id"] == ids["profile_version"]
    assert package["company_snapshot_at"]

    # FR-321: (b) and (c) are never sent.
    attachments = package_module.dispatch_attachments(package)
    assert [Path(a).name for a in attachments] == ["cv.pdf"]
    assert package["briefing_pdf_path"] not in attachments
    assert package["motivation_pdf_path"] not in attachments
    with pytest.raises(package_module.NeverSent):
        package_module.document_path(package, "briefing", for_dispatch=True)
    with pytest.raises(package_module.NeverSent):
        package_module.document_path(package, "motivation", for_dispatch=True)
    assert package_module.document_path(package, "cv_pdf").is_file()


def test_consistency_failure_blocks_approval_until_overridden() -> None:
    """FR-322: nothing is offered for dispatch before the check passes."""
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import package as package_module

    ids = seed()
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    notes = dict(package["generation_notes"])
    notes["cv_document"]["experience"][0]["company"] = "Globex International"
    repo.update_package(package["id"], ids["seeker"], {"generation_notes": notes})
    package = package_module.run_consistency(
        ids["seeker"], repo.get_package(package["id"], ids["seeker"])
    )
    assert package["consistency_status"] == "fail"

    blockers = package_module.approval_blockers(package)
    assert [b["kind"] for b in blockers] == ["consistency"]
    assert blockers[0]["overridable"] is True

    refused = package_module.approve(ids["seeker"], [package["id"]], actor="t")
    assert refused["approved"] == []

    accepted = package_module.approve(
        ids["seeker"], [package["id"]], actor="t", override_reason="Checked by hand"
    )
    assert accepted["approved"] == [package["id"]]
    assert repo.get_package(package["id"], ids["seeker"])["status"] == "approved"


def test_an_approved_override_survives_to_the_send_paths() -> None:
    """FR-322/FR-324: the recorded reason is on the row, not only in the audit.

    Approval offered an override and both send paths refused it, so an
    overridden package could never actually be sent (E2E_1500 section 8.7).
    """
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import package as package_module
    from dreamjob.documents.package import GenerationOptions

    ids = seed()
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    notes = dict(package["generation_notes"])
    notes["cv_document"]["experience"][0]["company"] = "Globex International"
    repo.update_package(package["id"], ids["seeker"], {"generation_notes": notes})
    package = package_module.run_consistency(
        ids["seeker"], repo.get_package(package["id"], ids["seeker"])
    )
    assert package["consistency_status"] == "fail"

    package_module.approve(
        ids["seeker"], [package["id"]], actor="t", override_reason="Checked by hand"
    )
    approved = repo.get_package(package["id"], ids["seeker"])
    assert approved["status"] == "approved"
    assert approved["consistency_override"] == "Checked by hand"

    # A regeneration is a fresh check, and the old reason does not answer it.
    package_module.generate(
        ids["seeker"],
        ids["opportunity"],
        GenerationOptions(instructions="redo"),
        package_id=package["id"],
        actor="t",
    )
    regenerated = repo.get_package(package["id"], ids["seeker"])
    assert regenerated["status"] == "draft"
    assert not regenerated.get("consistency_override")


def test_leak_failure_cannot_be_overridden() -> None:
    """NFR-206: a leak is never a judgement call."""
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import package as package_module

    ids = seed()
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    repo.update_package(package["id"], ids["seeker"], {"leak_scan_status": "fail"})
    package = repo.get_package(package["id"], ids["seeker"])

    blockers = package_module.approval_blockers(package)
    assert any(b["kind"] == "leak" and not b["overridable"] for b in blockers)
    result = package_module.approve(
        ids["seeker"], [package["id"]], actor="t", override_reason="I insist"
    )
    assert result["approved"] == []


def test_bulk_approval_requires_a_summary_and_records_it() -> None:
    """FR-324: bulk approval states what goes to whom, and it is audited."""
    from dreamjob.db.connection import query_all
    from dreamjob.documents import package as package_module

    ids = seed()
    first = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    second = package_module.generate(ids["seeker"], ids["speculative"], actor="test")
    package_ids = [first["id"], second["id"]]

    with pytest.raises(package_module.NotApprovable):
        package_module.approve(ids["seeker"], package_ids, actor="t")

    summary = package_module.bulk_summary(ids["seeker"], package_ids)
    assert summary["count"] == 2
    assert summary["recipients"] == ["els.peeters@zenith.example"]
    assert summary["blocked"] == []

    result = package_module.approve(
        ids["seeker"], package_ids, actor="tester", summary="Two applications to Zenith"
    )
    assert sorted(result["approved"]) == sorted(package_ids)

    events = query_all(
        "SELECT detail FROM audit_event WHERE action = 'application_package.approve'"
    )
    assert len(events) == 1
    assert "Two applications to Zenith" in events[0]["detail"]
    assert "els.peeters@zenith.example" in events[0]["detail"]


def test_editing_reopens_an_approved_package() -> None:
    """FR-324: an edit is a new draft, re-checked before it can be approved again."""
    from dreamjob.documents import package as package_module

    ids = seed()
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    package_module.approve(ids["seeker"], [package["id"]], actor="t")

    edited = package_module.edit(
        ids["seeker"], package["id"], {"email_body": "Beste Els,\n\nKorte vraag.\n"}
    )
    assert edited["status"] == "draft"
    assert edited["approved_at"] is None
    assert "Korte vraag" in edited["email_body"]


def test_briefing_refresh_survives_a_sent_package() -> None:
    """FR-329: refreshed just before an interview, which is after dispatch."""
    from dreamjob.db.repositories import applications as repo
    from dreamjob.documents import package as package_module

    ids = seed()
    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    repo.update_package(package["id"], ids["seeker"], {"status": "sent"})

    refreshed = package_module.generate(
        ids["seeker"], ids["opportunity"],
        package_module.GenerationOptions(parts=("briefing",)),
        package_id=package["id"], actor="test",
    )
    assert refreshed["status"] == "sent", "a briefing refresh must not withdraw a dispatch"
    assert Path(refreshed["briefing_pdf_path"]).is_file()

    with pytest.raises(package_module.GenerationError):
        package_module.generate(
            ids["seeker"], ids["opportunity"],
            package_module.GenerationOptions(parts=("cv", "email")),
            package_id=package["id"], actor="test",
        )


def test_retemplating_keeps_the_facts(tmp_path: Path) -> None:
    """FR-322: another template is a different layout, not a regeneration."""
    from dreamjob.documents import package as package_module
    from dreamjob.documents.templates import from_dict

    ids = seed(photo_path=_photo(tmp_path))
    package = package_module.generate(
        ids["seeker"], ids["opportunity"], package_module.GenerationOptions(cv_template="classic"),
        actor="test",
    )
    before = from_dict(package["generation_notes"]["cv_document"])

    updated = package_module.regenerate_cv_only(ids["seeker"], package["id"], template="modern")
    after = from_dict(updated["generation_notes"]["cv_document"])

    assert updated["cv_template"] == "modern"
    assert Path(updated["cv_pdf_path"]).is_file()
    assert [j.company for j in after.experience] == [j.company for j in before.experience]
    assert after.skills == before.skills


def test_generation_degrades_without_consent(tmp_path: Path) -> None:
    """CR-410: with no transfer consent recorded, no profile data reaches the API."""
    from dreamjob.documents import package as package_module

    ids = seed()
    client, reason = package_module.llm_for(ids["seeker"], None)
    assert client is None
    assert "consent" in reason.lower()

    package = package_module.generate(ids["seeker"], ids["opportunity"], actor="test")
    degradation = " ".join(package["generation_notes"].get("degradation") or [])
    assert "consent" in degradation.lower()


# ---------------------------------------------------------------------------
# The API surface (FR-324, FR-331, FR-344)
# ---------------------------------------------------------------------------


def _client(seeker_id: str, email: str = "stephane@stepvda.com") -> TestClient:
    """Only this slice's router, with authentication stubbed to one job seeker."""
    from dreamjob.api.deps import CurrentSeeker, current_seeker
    from dreamjob.api.routers import applications as router_module

    app = FastAPI()
    app.include_router(router_module.router, prefix="/api/applications")
    app.dependency_overrides[current_seeker] = lambda: CurrentSeeker(
        id=seeker_id, email=email, display_name="Test", is_admin=False, locale="nl"
    )
    return TestClient(app)


def test_api_generate_preview_and_download() -> None:
    ids = seed()
    client = _client(ids["seeker"])

    templates = client.get("/api/applications/templates").json()
    assert {t["key"] for t in templates["cv"]} >= {"classic", "modern"}
    assert templates["never_sent"] == ["briefing", "motivation"]

    # background=False asks for the synchronous form; the browser asks for the
    # job (see the assertion below), because minutes of model calls behind a
    # spinner read as a hang.
    created = client.post(
        "/api/applications/generate",
        json={
            "opportunity_ids": [ids["opportunity"]],
            "cv_template": "classic",
            "background": False,
        },
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["mode"] == "inline"
    package_id = body["package"]["id"]
    assert body["package"]["attachments"] == ["cv.pdf"]

    assert client.get(f"/api/applications/{package_id}").status_code == 200
    resolved = client.get(f"/api/applications/by-opportunity/{ids['opportunity']}").json()
    assert resolved["id"] == package_id

    for kind, media in (("cv_pdf", "application/pdf"), ("briefing", "application/pdf")):
        response = client.get(f"/api/applications/{package_id}/documents/{kind}")
        assert response.status_code == 200, kind
        assert response.headers["content-type"] == media
    assert client.get(f"/api/applications/{package_id}/documents/nope").status_code == 404

    edited = client.patch(
        f"/api/applications/{package_id}", json={"email_subject": "Sollicitatie Zenith"}
    )
    assert edited.status_code == 200
    assert edited.json()["email_subject"] == "Sollicitatie Zenith"

    listing = client.get("/api/applications/").json()
    assert listing["counts"] == {"draft": 1}


def test_api_bulk_approval_requires_a_summary() -> None:
    ids = seed()
    client = _client(ids["seeker"])
    packages = [
        client.post(
            "/api/applications/generate",
            json={"opportunity_ids": [opportunity], "background": False},
        ).json()["package"]["id"]
        for opportunity in (ids["opportunity"], ids["speculative"])
    ]

    refused = client.post("/api/applications/approve", json={"package_ids": packages})
    assert refused.status_code == 400
    assert "summary" in refused.json()["detail"].lower()

    summary = client.post(
        "/api/applications/approval-summary", json={"package_ids": packages}
    ).json()
    assert summary["count"] == 2
    assert summary["recipients"] == ["els.peeters@zenith.example"]

    approved = client.post(
        "/api/applications/approve",
        json={"package_ids": packages, "summary": "Two applications to Zenith"},
    )
    assert approved.status_code == 200
    assert sorted(approved.json()["approved"]) == sorted(packages)


def test_api_refuses_another_seekers_package() -> None:
    """FR-344: a package is only ever reachable by the job seeker who owns it."""
    from dreamjob.db.connection import insert_row, utcnow

    ids = seed()
    package_id = (
        _client(ids["seeker"])
        .post(
            "/api/applications/generate",
            json={"opportunity_ids": [ids["opportunity"]], "background": False},
        )
        .json()["package"]["id"]
    )
    intruder = insert_row(
        "job_seeker",
        {"email": "other@example.com", "display_name": "Other", "locale": "en",
         "created_at": utcnow(), "updated_at": utcnow()},
    )
    other = _client(intruder, "other@example.com")

    assert other.get(f"/api/applications/{package_id}").status_code == 404
    assert other.get(f"/api/applications/{package_id}/documents/cv_pdf").status_code == 404
    assert other.get("/api/applications/").json()["packages"] == []
