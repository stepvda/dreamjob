"""End-to-end: a fictional IT profile through autopilot to application documents.

This is the honest test of the three-step promise. It creates a persona that
has never existed, gives it a CV and a dream-job statement, starts the
autopilot, waits for the shortlist, selects an opportunity and generates the
application package (CV, briefing, motivation, email) — stopping before send.

Run against the live API:
    PYTHONPATH=backend python3 -m scripts.e2e_persona          # full run
"""

from __future__ import annotations

import io
import json
import os
import sys
import time
import urllib.error
import urllib.request

API = os.environ.get("DREAMJOB_API", "http://127.0.0.1:8000")
PASSWORD = "Persona-Str0ng!Pass"
EMAIL = os.environ.get("PERSONA_EMAIL", f"alex.dev+persona-{int(time.time())}@example.com")
ADMIN_EMAIL = "stephane@stepvda.com"
ADMIN_PASSWORD = "Adv39281#"

CV_TEXT = [
    ("Alex De Vries", "Senior Backend & Platform Engineer"),
    ("Brussels, Belgium | alex.dev@example.com | github.com/alexdev"),
    ("", ""),
    ("SUMMARY", "Backend engineer, 9 years, Python and Go, distributed systems and APIs."),
    ("", ""),
    ("EXPERIENCE", ""),
    ("Senior Backend Engineer", "MediCloud — Brussels  2021 - Present"),
    ("- Led the migration of a patient-data API from a monolith to FastAPI services."),
    ("- Built an event pipeline in Kafka handling 40M events/day; cut p95 latency 60%."),
    ("- Owned GDPR data-subject erasure across six services."),
    ("Backend Engineer", "Fintech BV — Ghent  2018 - 2021"),
    ("- Designed a payments reconciliation service in Python and PostgreSQL."),
    ("- Introduced CI/CD and infrastructure-as-code with Terraform on AWS."),
    ("Software Engineer", "Studio Zes — Antwerp  2016 - 2018"),
    ("- Built REST APIs and internal tooling; first production Kubernetes deployment."),
    ("", ""),
    ("SKILLS", "Python, FastAPI, Go, PostgreSQL, Kafka, Docker, Kubernetes, AWS, Terraform, CI/CD, REST, GraphQL, GDPR"),
    ("", ""),
    ("EDUCATION", "MSc Computer Science — KU Leuven  2014 - 2016"),
    ("LANGUAGES", "Dutch (native), English (fluent), French (fluent)"),
]

DREAM_JOB = (
    "I want to keep building backend systems hands-on, in Python, at a place where "
    "engineering quality is taken seriously rather than talked about. I like small teams "
    "where I own a service end to end — design, build, operate. A product with real scale "
    "and a reason to exist: health, climate, public infrastructure, developer tools. "
    "I want to avoid pure management, and I do not want a role where I only maintain "
    "something someone else designed. Brussels or remote, no relocation."
)


def _request(method: str, path: str, *, token=None, body=None, files=None, raw=False):
    url = f"{API}{path}"
    headers = {}
    data = None
    if files is not None:
        boundary = "----dreamjobE2E"
        parts = []
        for name, (filename, content, ctype) in files.items():
            parts.append(
                f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"; "
                f"filename=\"{filename}\"\r\nContent-Type: {ctype}\r\n\r\n".encode()
                + content
                + b"\r\n"
            )
        parts.append(f"--{boundary}--\r\n".encode())
        data = b"".join(parts)
        headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    elif body is not None:
        data = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    if token:
        headers["Authorization"] = f"Bearer {token}"

    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            payload = resp.read()
            return resp.status, (payload if raw else json.loads(payload or b"null"))
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read().decode()[:300]


def _make_cv_docx() -> bytes:
    """A real DOCX so the CV parser is exercised, not bypassed."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for line in CV_TEXT:
        if isinstance(line, tuple):
            left, right = line
            p = doc.add_paragraph()
            run = p.add_run(left)
            run.bold = True
            if right:
                run2 = p.add_run(f"   {right}")
                run2.font.size = Pt(10)
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def step(msg: str) -> None:
    print(f"\n=== {msg} ===", flush=True)


def main() -> int:
    step(f"Persona: {EMAIL}")
    status, body = _request("POST", "/api/auth/register", body={
        "email": EMAIL, "display_name": "Alex De Vries", "password": PASSWORD,
    })
    if status not in (200, 201):
        print("register:", status, body)
        return 1
    token = body["session"]["token"]
    print("registered, admin:", body["is_admin"])

    _request("POST", "/api/auth/consent", token=token,
             body={"kind": "llm_transfer", "granted": True})

    step("Upload CV")
    cv = _make_cv_docx()
    status, body = _request(
        "POST", "/api/profile/uploads/cv", token=token,
        files={"file": ("alex-cv.docx", cv,
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    print("cv upload:", status, "" if status < 400 else body)

    step("Dream-job statement")
    status, body = _request("PUT", "/api/profile/dream-job", token=token,
                            body={"statement": DREAM_JOB})
    print("statement:", status, "" if status < 400 else body)

    step("Preflight")
    status, pre = _request("GET", "/api/autopilot/preflight", token=token)
    print(json.dumps(pre, indent=1)[:400] if isinstance(pre, dict) else pre)

    step("Start autopilot (bounded)")
    status, body = _request("POST", "/api/autopilot/start", token=token,
                            body={"company_limit": 6})
    print("start:", status, body if status >= 400 else body.get("job_id"))
    if status >= 400:
        return 1

    step("Waiting for the shortlist")
    deadline = time.time() + 45 * 60
    last = None
    while time.time() < deadline:
        status, st = _request("GET", "/api/autopilot/status", token=token)
        run = (st or {}).get("run") or {}
        marker = (run.get("status"), run.get("stage"), run.get("stage_index"))
        if marker != last:
            print(f"  {run.get('status'):9} {str(run.get('stage')):12} "
                  f"{run.get('stage_index')}/{run.get('total_steps')} "
                  f"errs={run.get('error_count')}", flush=True)
            last = marker
        if run.get("status") in ("done", "failed", "cancelled"):
            print("\nfinal report:", json.dumps(run.get("report"), indent=1)[:900])
            print("last_error:", run.get("last_error"))
            break
        time.sleep(3)
    else:
        print("timed out waiting for autopilot")
        return 1

    step("Shortlist")
    status, opps = _request("GET", "/api/opportunities?limit=10", token=token)
    items = opps if isinstance(opps, list) else opps.get("items") or opps.get("opportunities") or []
    print("opportunities:", len(items))
    for o in items[:8]:
        print(f"  [{str(o.get('kind')):11}] {(o.get('title') or '')[:44]:44} "
              f"{str(o.get('company_name') or '')[:22]:22} score={o.get('score')}")
    if not items:
        print("no opportunities to apply to")
        return 1

    step("Generate the application package for the top opportunity")
    top = items[0]
    # Contacts first so the email is addressed to a person where one can be
    # found; best-effort, because a campaign can have none.
    status, contacts = _request("POST", "/api/apply/contacts/discover", token=token,
                                body={"opportunity_ids": [top["id"]]})
    print("contact discovery:", status)

    status, body = _request("POST", "/api/apply/packages/generate", token=token,
                            body={"opportunity_ids": [top["id"]]})
    print("generate:", status, body if status >= 400 else json.dumps(body)[:300])
    if status >= 400:
        return 1

    step("Waiting for the package")
    package = None
    for _ in range(300):
        status, view = _request("GET", f"/api/apply/{top['id']}", token=token)
        if status < 400 and isinstance(view, dict):
            pkg = view.get("package") or view
            if pkg.get("id") and pkg.get("status") not in (None, "generating", "queued"):
                package = view
                break
        time.sleep(2)

    step("Application package")
    if not package:
        print("no package appeared within the wait")
        return 1
    pkg = package.get("package") or package
    for key in (
        "status", "language", "cv_docx_path", "cv_pdf_path", "briefing_pdf_path",
        "motivation_pdf_path", "consistency_status", "leak_scan_status",
        "email_subject",
    ):
        print(f"  {key}: {pkg.get(key)}")
    body_text = pkg.get("email_body") or ""
    print("\n  email (first 500 chars):")
    print("  " + body_text[:500].replace("\n", "\n  "))

    step("Artefacts on disk")
    import pathlib

    root = pathlib.Path(__file__).resolve().parents[1]
    for key in ("cv_docx_path", "cv_pdf_path", "briefing_pdf_path", "motivation_pdf_path"):
        raw = pkg.get(key)
        if not raw:
            continue
        p = pathlib.Path(raw)
        if not p.is_absolute():
            p = root / raw
        print(f"  {key}: {'OK' if p.exists() else 'MISSING'}  {raw}  "
              f"({p.stat().st_size if p.exists() else 0} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
