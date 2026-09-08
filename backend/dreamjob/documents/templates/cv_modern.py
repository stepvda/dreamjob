"""Modern CV template: coloured sidebar plus main column (FR-322).

The same ``CvDocument`` as the classic template, laid out as a two-column page:
photograph, contact details, skills and languages in a tinted sidebar, the
narrative - profile, experience, education - in the main column.

The PDF uses two frames on the first page and a single full-width frame on any
continuation page, rather than one long two-column table.  A CV that runs to a
second page then keeps reading as a document instead of stranding a column of
skills opposite an empty half-page.
"""

from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    FrameBreak,
    Image,
    NextPageTemplate,
    PageTemplate,
    Paragraph,
    Spacer,
)

from dreamjob.documents.pdf_builder import (
    HorizontalRule,
    PdfBuilder,
    escape,
    footer_canvas,
)
from dreamjob.documents.templates import CvDocument
from dreamjob.documents.templates._docx import (
    borderless,
    bullet,
    cell_margins,
    new_document,
    para,
    rule,
    shade,
)

TEMPLATE_KEY = "modern"
DISPLAY_NAME = "Modern"
DESCRIPTION = (
    "Two columns: a tinted sidebar with photograph, contact details, skills and "
    "languages, and a main column with the narrative."
)

SIDEBAR_MM = 58.0
MARGIN_MM = 14.0
SIDEBAR_TINT = colors.HexColor("#eef3f7")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def render_docx(cv: CvDocument, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = new_document(margin_cm=1.4)
    accent = cv.accent

    table = document.add_table(rows=1, cols=2)
    borderless(table)
    sidebar, main = table.rows[0].cells
    sidebar.width = _cm(5.6)
    main.width = _cm(12.2)
    shade(sidebar, "#eef3f7")
    cell_margins(sidebar, left=0.35, right=0.35, top=0.3, bottom=0.3)
    cell_margins(main, left=0.5, right=0.0, top=0.3)

    _docx_sidebar(sidebar, cv, accent)
    _docx_main(main, cv, accent)

    _docx_footer(document, cv)
    document.save(str(target))
    return target


def _docx_sidebar(cell, cv: CvDocument, accent: str) -> None:
    first = cell.paragraphs[0]
    first.paragraph_format.space_after = _pt(6)
    if cv.has_photo():
        first.alignment = 1  # centre
        first.add_run().add_picture(str(cv.photo_path), width=_cm(3.4))

    contact_lines = cv.contact.lines()
    if contact_lines:
        _docx_sidebar_heading(cell, cv.label("cv_curriculum_vitae"), accent)
        for line in contact_lines:
            para(cell, line, size=8.5, space_after=1)

    if cv.skills:
        _docx_sidebar_heading(cell, cv.section_label("skills"), accent)
        for skill in cv.skills:
            para(cell, skill, size=8.5, space_after=1)

    if cv.languages:
        _docx_sidebar_heading(cell, cv.section_label("languages"), accent)
        for entry in cv.languages:
            para(cell, entry.line(), size=8.5, space_after=1)

    for key, entries in cv.entry_sections():
        _docx_sidebar_heading(cell, cv.section_label(key), accent)
        for entry in entries:
            text = f"{entry.title} ({entry.date})" if entry.date else entry.title
            para(cell, text, size=8.5, space_after=1)


def _docx_main(cell, cv: CvDocument, accent: str) -> None:
    first = cell.paragraphs[0]
    first.paragraph_format.space_after = _pt(0)
    run = first.add_run(cv.contact.name)
    run.bold = True
    run.font.size = _pt(22)
    run.font.color.rgb = _rgb(accent)
    if cv.contact.headline:
        para(cell, cv.contact.headline, size=11, colour="#5b6b7c", space_after=4)
    rule(para(cell, space_after=6), accent, size=12)

    if cv.summary:
        _docx_heading(cell, cv.section_label("summary"), accent)
        para(cell, cv.summary, size=9.5, align="justify")

    if cv.experience:
        _docx_heading(cell, cv.section_label("experience"), accent)
        for job in cv.experience:
            heading = para(cell, space_before=4, space_after=0)
            run = heading.add_run(job.title or job.company)
            run.bold = True
            run.font.size = _pt(10.5)
            if job.company and job.title:
                para(cell, job.company, size=9.5, colour=accent, space_after=0)
            meta = " · ".join(p for p in (job.period(cv.language), job.location) if p)
            if meta:
                para(cell, meta, size=8.5, italic=True, colour="#5b6b7c", space_after=2)
            for line in job.bullets:
                bullet(cell, line, size=9.5)

    if cv.education:
        _docx_heading(cell, cv.section_label("education"), accent)
        for study in cv.education:
            heading = para(cell, space_before=3, space_after=0)
            run = heading.add_run(study.school)
            run.bold = True
            run.font.size = _pt(10)
            line = " · ".join(p for p in (study.line(), study.period(cv.language)) if p)
            if line:
                para(cell, line, size=8.5, italic=True, colour="#5b6b7c", space_after=2)
            if study.detail:
                para(cell, study.detail, size=9)


def _docx_heading(container, text: str, accent: str) -> None:
    paragraph = para(
        container, text.upper(), size=10, bold=True, colour=accent,
        space_before=10, space_after=2,
    )
    rule(paragraph, accent)


def _docx_sidebar_heading(container, text: str, accent: str) -> None:
    para(container, text.upper(), size=8.5, bold=True, colour=accent,
         space_before=8, space_after=2)


def _docx_footer(document, cv: CvDocument) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.text = cv.meta.footer_line()
    for run in paragraph.runs:
        run.font.size = _pt(7)


def _pt(value: float):  # noqa: ANN202
    from docx.shared import Pt  # noqa: PLC0415

    return Pt(value)


def _cm(value: float):  # noqa: ANN202
    from docx.shared import Cm  # noqa: PLC0415

    return Cm(value)


def _rgb(value: str):  # noqa: ANN202
    from dreamjob.documents.templates._docx import hex_to_rgb  # noqa: PLC0415

    return hex_to_rgb(value)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def render_pdf(cv: CvDocument, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    accent = colors.HexColor(cv.accent)

    sidebar = PdfBuilder(cv.meta, accent=accent)
    main = PdfBuilder(cv.meta, accent=accent)
    _pdf_sidebar(sidebar, cv)
    _pdf_main(main, cv)

    page_width, page_height = A4
    margin = MARGIN_MM * mm
    sidebar_width = SIDEBAR_MM * mm
    gutter = 6 * mm
    body_bottom = margin + 8 * mm
    body_height = page_height - margin - body_bottom

    sidebar_frame = Frame(
        margin, body_bottom, sidebar_width - 2 * mm, body_height,
        leftPadding=4, rightPadding=6, topPadding=6, bottomPadding=6, id="sidebar",
    )
    main_frame = Frame(
        margin + sidebar_width + gutter, body_bottom,
        page_width - 2 * margin - sidebar_width - gutter, body_height,
        leftPadding=0, rightPadding=0, topPadding=6, bottomPadding=6, id="main",
    )
    full_frame = Frame(
        margin, body_bottom, page_width - 2 * margin, body_height,
        leftPadding=0, rightPadding=0, topPadding=6, bottomPadding=6, id="full",
    )

    def paint_sidebar(canvas, _doc) -> None:
        canvas.saveState()
        canvas.setFillColor(SIDEBAR_TINT)
        canvas.rect(margin, body_bottom, sidebar_width, body_height, stroke=0, fill=1)
        canvas.setFillColor(accent)
        canvas.rect(margin, body_bottom + body_height - 3 * mm, sidebar_width, 3 * mm,
                    stroke=0, fill=1)
        canvas.restoreState()

    doc = BaseDocTemplate(
        str(target),
        pagesize=A4,
        title=cv.meta.title,
        author=cv.contact.name or "Dream Job",
        subject=cv.meta.subtitle,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=body_bottom,
    )
    doc.addPageTemplates(
        [
            PageTemplate(id="two-column", frames=[sidebar_frame, main_frame],
                         onPage=paint_sidebar),
            PageTemplate(id="single-column", frames=[full_frame]),
        ]
    )

    story = [
        *sidebar.story,
        FrameBreak(),
        NextPageTemplate("single-column"),
        *main.story,
    ]
    doc.build(story, canvasmaker=footer_canvas(cv.meta))
    return target


def _pdf_sidebar(sidebar: PdfBuilder, cv: CvDocument) -> None:
    styles = sidebar.styles
    if cv.has_photo():
        width = (SIDEBAR_MM - 16) * mm
        sidebar.story.append(
            Image(str(cv.photo_path), width=width, height=width, kind="proportional")
        )
        sidebar.spacer(4)
    for line in cv.contact.lines():
        sidebar.story.append(Paragraph(escape(line), styles["small"]))
    if cv.skills:
        _sidebar_heading(sidebar, cv.section_label("skills"))
        for skill in cv.skills:
            sidebar.story.append(Paragraph(escape(skill), styles["small"]))
    if cv.languages:
        _sidebar_heading(sidebar, cv.section_label("languages"))
        for entry in cv.languages:
            sidebar.story.append(Paragraph(escape(entry.line()), styles["small"]))
    for key, entries in cv.entry_sections():
        _sidebar_heading(sidebar, cv.section_label(key))
        for entry in entries:
            text = f"{entry.title} ({entry.date})" if entry.date else entry.title
            sidebar.story.append(Paragraph(escape(text), styles["small"]))


def _sidebar_heading(sidebar: PdfBuilder, text: str) -> None:
    sidebar.spacer(3)
    sidebar.story.append(Paragraph(f"<b>{escape(text.upper())}</b>", sidebar.styles["h2"]))
    sidebar.story.append(HorizontalRule(colour=sidebar.accent, thickness=0.7))


def _pdf_main(main: PdfBuilder, cv: CvDocument) -> None:
    main.story.append(Paragraph(escape(cv.contact.name), main.styles["title"]))
    if cv.contact.headline:
        main.story.append(Paragraph(escape(cv.contact.headline), main.styles["subtitle"]))
    main.story.append(HorizontalRule(colour=main.accent, thickness=1.2))
    main.spacer(2)

    if cv.summary:
        main.h1(cv.section_label("summary"))
        main.para(cv.summary)

    if cv.experience:
        main.h1(cv.section_label("experience"))
        for job in cv.experience:
            main.keep_together(lambda sub, job=job: _pdf_job(sub, cv, job))

    if cv.education:
        main.h1(cv.section_label("education"))
        for study in cv.education:
            main.h2(study.school)
            line = " · ".join(p for p in (study.line(), study.period(cv.language)) if p)
            if line:
                main.para(line, style="note")
            if study.detail:
                main.para(study.detail)


def _pdf_job(sub: PdfBuilder, cv: CvDocument, job) -> None:
    sub.story.append(Paragraph(f"<b>{escape(job.title or job.company)}</b>", sub.styles["h2"]))
    if job.company and job.title:
        sub.story.append(
            Paragraph(
                f'<font color="{cv.accent}">{escape(job.company)}</font>', sub.styles["bullet"]
            )
        )
    meta = " · ".join(p for p in (job.period(cv.language), job.location) if p)
    if meta:
        sub.story.append(Paragraph(escape(meta), sub.styles["note"]))
    sub.bullets(job.bullets)
    sub.story.append(Spacer(1, 2 * mm))
