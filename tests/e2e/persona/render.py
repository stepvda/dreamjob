#!/usr/bin/env python3
"""Render the persona into the two documents the product ingests (FR-102, FR-103).

``generate.py`` writes the persona; this turns it into the pair of files a job
seeker actually uploads:

* ``out/persona_cv.docx`` - a CV in the shape the CV parser was written for: a
  table-layout contact header, upper-case bold section headings, an entry line
  per role with the dates behind a tab, achievement bullets, and a photograph
  in ``word/media`` so the FR-103 photo-extraction path has something to find.
* ``out/persona_linkedin.pdf`` - a LinkedIn "Save to PDF" export.

The export is the fussier of the two.  ``linkedin_pdf`` does not read the text
stream; it reads word boxes, splitting the rail from the body on ``x0`` and
telling a section header from a company from a role title from a date line on
font size alone.  Every measurement below was taken from a real export
(``docs/Profile.pdf``): a 612x792 page, the rail at x=21.6 and the body at
x=223.6, the name at 26pt, body headers at 15.75, rail headers at 13, company
and headline at 12, role titles at 11.5 and everything else at 10.5.  A fixture
that merely looked like a LinkedIn export would prove nothing, so this one is
built to those numbers and then parsed back to prove it.

The photograph is an abstract monogram drawn with PIL - a coloured geometric
mark, deliberately not a synthetic face: the test needs bytes travelling
through the extraction path, not an invented likeness of a person who does not
exist.

Both outputs are deterministic given the same ``persona.json``: the palette is
seeded from the name, the PDF is written with reportlab's ``invariant`` flag,
and the DOCX timestamps are pinned to the persona's own generation date.

    python3 tests/e2e/persona/render.py
    python3 tests/e2e/persona/render.py --persona out/persona.json

The run finishes by parsing both documents back with the product's own
parsers, merging them, and checking that the two planted disagreements turned
into conflict rows.  A fixture whose conflicts the merge cannot see would let
the end-to-end test pass while proving nothing, so that check sets the exit
code.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas as pdfcanvas

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT / "backend"))

from dreamjob.pipeline.cv_parser import parse_cv  # noqa: E402
from dreamjob.pipeline.linkedin_pdf import parse_linkedin_pdf  # noqa: E402
from dreamjob.pipeline.profile_intake import merge_documents  # noqa: E402

OUT_DIR = Path(__file__).resolve().parent / "out"
PERSONA_PATH = OUT_DIR / "persona.json"
CV_PATH = OUT_DIR / "persona_cv.docx"
LINKEDIN_PATH = OUT_DIR / "persona_linkedin.pdf"
AVATAR_PATH = OUT_DIR / "persona_avatar.png"

# ---------------------------------------------------------------------------
# LinkedIn export geometry, measured from docs/Profile.pdf
# ---------------------------------------------------------------------------

PAGE_W, PAGE_H = 612.0, 792.0
RAIL_X, RAIL_W = 21.6, 186.0
BODY_X, BODY_W = 223.6, 352.0
TOP_MARGIN = 44.0
BODY_BOTTOM = 748.0
LEADING = 1.45
FONT = "Helvetica"

SIZE_NAME = 26.0
SIZE_BODY_HEADER = 15.75
SIZE_RAIL_HEADER = 13.0
SIZE_ENTITY = 12.0       # headline, location, company name, school name
SIZE_ROLE = 11.5
SIZE_DETAIL = 10.5
SIZE_RAIL_URL = 11.0
SIZE_FOOTER = 9.0
FOOTER_X, FOOTER_Y = 384.0, 772.0

# ---------------------------------------------------------------------------
# Avatar
# ---------------------------------------------------------------------------

AVATAR_PX = 512

# Deep ground, mid figure, warm accent - readable once the DOCX scales it to a
# thumbnail, and nothing that could be mistaken for a photograph of a person.
_PALETTES: tuple[tuple[str, str, str], ...] = (
    ("#16324F", "#3E7CB1", "#F4B860"),
    ("#2D3142", "#4F5D75", "#EF8354"),
    ("#14342B", "#3B8C6E", "#C6E5B1"),
    ("#2B1B4A", "#7A5AA8", "#E4B7E5"),
    ("#0B2545", "#356390", "#8DA9C4"),
)

_FONT_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
)


def _monogram_font(size: int) -> Any:
    for candidate in _FONT_CANDIDATES:
        path = Path(candidate)
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size)
            except OSError:  # a font file that will not load is not a failure
                continue
    return ImageFont.load_default(size=size)


def build_avatar(persona: dict[str, Any], path: Path) -> Path:
    """Draw a coloured geometric monogram, seeded from the name.

    A synthetic face would be both harder to justify and no better as a test:
    what the intake path cares about is that the DOCX carries an image part it
    can pull out and store (FR-103).
    """
    identity = persona["identity"]
    initials = f"{identity['first_name'][:1]}{identity['last_name'][:1]}".upper()
    seed = hashlib.sha256(identity["name"].encode("utf-8")).digest()
    ground, figure, accent = _PALETTES[seed[0] % len(_PALETTES)]

    size = AVATAR_PX
    image = Image.new("RGB", (size, size), ground)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(
        (size * 0.05, size * 0.05, size * 0.95, size * 0.95),
        radius=int(size * 0.12),
        fill=figure,
    )
    # Two flat shapes, placed from the seed so two personas never look alike.
    corner = seed[1] % 4
    span = size * 0.62
    origins = ((0.0, 0.0), (size - span, 0.0), (0.0, size - span), (size - span, size - span))
    ox, oy = origins[corner]
    draw.polygon([(ox, oy), (ox + span, oy), (ox, oy + span)], fill=ground)
    radius = size * (0.16 + (seed[2] % 5) * 0.01)
    cx, cy = size * 0.74, size * 0.74
    draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=accent)

    draw.text(
        (size / 2, size / 2),
        initials,
        font=_monogram_font(int(size * 0.42)),
        fill="#FFFFFF",
        anchor="mm",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=True)
    return path


# ---------------------------------------------------------------------------
# CV (DOCX)
# ---------------------------------------------------------------------------

_HEADING_RGB = RGBColor(0x16, 0x32, 0x4F)
_MUTED_RGB = RGBColor(0x5A, 0x63, 0x72)
_TAB_CM = 17.4      # the right edge of the text column, where the dates sit


def _bottom_rule(paragraph: Any) -> None:
    """Hairline under a section heading.

    python-docx has no API for paragraph borders, so this reaches into the
    element the way the library's own documentation does.
    """
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "B6BECA")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _heading(document: Any, text: str) -> None:
    """An upper-case, fully bold, short line - what the CV parser reads as a heading."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = _HEADING_RGB
    _bottom_rule(paragraph)


def _entry(document: Any, lead: str, rest: str, dates: str | None = None) -> None:
    """One entry line: a bold lead, middot-separated fields, dates behind a tab.

    The padded middot is what ``cv_parser._split_fields`` splits on, and the tab
    is what ``_split_dates`` looks for first; both are also how a CV of this
    shape is actually typed.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    if dates:
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(_TAB_CM), WD_TAB_ALIGNMENT.RIGHT
        )
    run = paragraph.add_run(lead)
    run.bold = True
    if rest:
        paragraph.add_run(rest)
    if dates:
        run = paragraph.add_run(f"\t{dates}")
        run.font.color.rgb = _MUTED_RGB


def _bullet(document: Any, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run(f"• {text}")


def _labelled(document: Any, label: str, body: str) -> None:
    """A bold category and its members - the CV parser's categorised skill line."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    paragraph.add_run(body)


def _contact_cell(cell: Any, persona: dict[str, Any]) -> None:
    """The contact block, in the order ``cv_parser._read_contact`` reads it.

    Line 0 is the name and line 1 the headline; everything the merge compares
    across the two sources (name, email, phone) is printed exactly as the
    export prints it, so the only disagreements are the planted ones.
    """
    identity = persona["identity"]
    lines = [
        (identity["name"], 18, True, _HEADING_RGB),
        (identity["headline"], 9.5, False, _MUTED_RGB),
        (identity["location_full"], 9.5, False, None),
        (identity["email"], 9.5, False, None),
        (identity["phone"], 9.5, False, None),
        (identity["linkedin_url"], 9.5, False, None),
        (identity["website"], 9.5, False, None),
    ]
    first = True
    for text, size, bold, colour in lines:
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if colour is not None:
            run.font.color.rgb = colour


def build_cv(persona: dict[str, Any], path: Path, photo: Path) -> Path:
    document = docx.Document()

    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.6)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.1

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(13.4), Cm(4.0)
    _contact_cell(table.cell(0, 0), persona)
    table.cell(0, 1).paragraphs[0].add_run().add_picture(str(photo), width=Cm(3.2))

    _heading(document, "Profile")
    document.add_paragraph(persona["summary"])

    _heading(document, "Experience")
    for role in persona["experience"]:
        _entry(
            document,
            role["title"],
            f"  ·  {role['employer_cv']}  ·  {role['location']}",
            role["cv_dates"],
        )
        _bullet(document, role["summary"])
        for achievement in role["achievements"]:
            _bullet(document, achievement)

    _heading(document, "Education")
    for study in persona["education"]:
        _entry(
            document,
            study["degree"],
            f"  ·  {study['school']}  ·  {study['location']}",
            study["cv_dates"],
        )
        _bullet(document, f"{study['field']}.")

    _heading(document, "Technical skills")
    for group in persona["skill_groups"]:
        _labelled(document, group["label"], ", ".join(group["skills"]))

    _heading(document, "Certifications")
    for certification in persona["certifications"]:
        _entry(
            document,
            certification["title"],
            f"  ·  {certification['issuer']}",
            str(certification["year"]),
        )

    _heading(document, "Languages")
    document.add_paragraph(
        ", ".join(f"{lang['language']} ({lang['proficiency']})" for lang in persona["languages"])
    )

    _heading(document, "Publications")
    for publication in persona["publications"]:
        _entry(
            document,
            publication["title"],
            f"  ·  {publication['venue']}",
            str(publication["year"]),
        )

    _heading(document, "Projects")
    for project in persona["projects"]:
        _entry(document, project["name"], f"  ·  {project['url']}")
        _bullet(document, project["description"])

    # Pinned so a re-run of the same persona produces the same document rather
    # than one that differs only in when it was written.
    stamp = datetime.fromisoformat(persona["generated_at"]).replace(tzinfo=UTC)
    properties = document.core_properties
    properties.author = persona["identity"]["name"]
    properties.title = f"CV - {persona['identity']['name']}"
    properties.comments = persona["notice"]
    properties.created = properties.modified = stamp
    properties.last_modified_by = persona["identity"]["name"]

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    _pin_zip_timestamps(path, stamp)
    return path


def _pin_zip_timestamps(path: Path, stamp: datetime) -> None:
    """Rewrite the archive with fixed entry timestamps.

    A DOCX is a zip, and python-docx stamps every entry with the wall clock,
    so two renders of the same persona would differ in bytes while being the
    same document. Pinning them makes the fixture diffable and lets a test
    compare checksums.
    """
    date_time = (stamp.year, stamp.month, stamp.day, 0, 0, 0)
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, payload in entries:
            pinned = zipfile.ZipInfo(info.filename, date_time)
            pinned.compress_type = info.compress_type
            pinned.create_system = info.create_system
            pinned.external_attr = info.external_attr
            target.writestr(pinned, payload)


# ---------------------------------------------------------------------------
# LinkedIn export (PDF)
# ---------------------------------------------------------------------------


def _wrap(text: str, size: float, width: float) -> list[str]:
    """Greedy word wrap at a measured width, so no line runs into the margin."""
    words = str(text).split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if not current or pdfmetrics.stringWidth(trial, FONT, size) <= width:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def _fitted(text: str, size: float, width: float) -> float:
    """Shrink an unbreakable line (a URL) until it stays inside the rail.

    A rail word that reached past the column boundary would be read as body
    text and the two columns would interleave (``_column_split_x``).
    """
    while size > 8.5 and pdfmetrics.stringWidth(text, FONT, size) > width:
        size -= 0.5
    return size


def _rail_lines(persona: dict[str, Any]) -> list[tuple[str, float, float]]:
    """The narrow left column: Contact, Top Skills, Languages."""
    identity = persona["identity"]
    items: list[tuple[str, float, float]] = [("Contact", SIZE_RAIL_HEADER, 0.0)]
    items.append((f"{identity['phone']} (Mobile)", SIZE_DETAIL, 8.0))
    items.append((identity["email"], SIZE_DETAIL, 0.0))
    for url, label in ((identity["linkedin_url"], "LinkedIn"), (identity["website"], "Personal")):
        size = _fitted(url, SIZE_RAIL_URL, RAIL_W)
        items.append((url, size, 6.0))
        items.append((f"({label})", SIZE_RAIL_URL, 0.0))

    items.append(("Top Skills", SIZE_RAIL_HEADER, 18.0))
    for i, skill in enumerate(persona["top_skills"]):
        items.append((skill, SIZE_DETAIL, 8.0 if i == 0 else 0.0))

    items.append(("Languages", SIZE_RAIL_HEADER, 18.0))
    for i, language in enumerate(persona["languages"]):
        text = f"{language['language']} ({language['proficiency']})"
        items.append((text, _fitted(text, SIZE_DETAIL, RAIL_W), 8.0 if i == 0 else 0.0))
    return items


def _body_items(persona: dict[str, Any]) -> list[tuple[str, float, float, bool]]:
    """The wide right column, as ``(text, size, gap_before, keep_with_next)``."""
    identity = persona["identity"]
    items: list[tuple[str, float, float, bool]] = [(identity["name"], SIZE_NAME, 0.0, True)]
    for i, line in enumerate(_wrap(identity["headline"], SIZE_ENTITY, BODY_W)):
        items.append((line, SIZE_ENTITY, 6.0 if i == 0 else 0.0, True))
    items.append((identity["location_full"], SIZE_ENTITY, 0.0, False))

    items.append(("Summary", SIZE_BODY_HEADER, 22.0, True))
    for i, line in enumerate(_wrap(persona["summary"], SIZE_ENTITY, BODY_W)):
        items.append((line, SIZE_ENTITY, 8.0 if i == 0 else 0.0, i == 0))

    items.append(("Experience", SIZE_BODY_HEADER, 26.0, True))
    previous_employer: str | None = None
    for role in persona["experience"]:
        head: list[tuple[str, float, float, bool]] = []
        if role["employer"] != previous_employer:
            head.append((role["employer"], SIZE_ENTITY, 20.0, True))
            if role["employer_total_duration"]:
                head.append((role["employer_total_duration"], SIZE_DETAIL, 4.0, True))
            previous_employer = role["employer"]
        # A title under its own company heading sits close to it; a second role
        # at the same employer needs the wider gap the export uses to separate
        # it from the description above.
        head.append((role["title"], SIZE_ROLE, 10.0 if head else 18.0, True))
        head.append((role["linkedin_dates"], SIZE_DETAIL, 3.0, True))
        head.append((role["location"], SIZE_DETAIL, 3.0, False))
        items.extend(head)
        paragraphs = [role["summary"], *(f"- {a}" for a in role["achievements"])]
        for paragraph in paragraphs:
            for i, line in enumerate(_wrap(paragraph, SIZE_DETAIL, BODY_W)):
                items.append((line, SIZE_DETAIL, 8.0 if i == 0 else 0.0, False))

    items.append(("Education", SIZE_BODY_HEADER, 26.0, True))
    for study in persona["education"]:
        items.append((study["school"], SIZE_ENTITY, 16.0, True))
        detail = _wrap(study["linkedin_detail"], SIZE_DETAIL, BODY_W)
        for i, line in enumerate(detail):
            items.append((line, SIZE_DETAIL, 5.0 if i == 0 else 0.0, i < len(detail) - 1))
    return items


def _flow(items: list[tuple[str, float, float, bool]]) -> list[list[tuple[float, str, float]]]:
    """Break the body into pages, never splitting a company from its role."""
    pages: list[list[tuple[float, str, float]]] = []
    page: list[tuple[float, str, float]] = []
    y = TOP_MARGIN
    index = 0
    while index < len(items):
        last = index
        while items[last][3] and last + 1 < len(items):
            last += 1
        run = items[index : last + 1]
        height = sum(gap + size * LEADING for _, size, gap, _ in run)
        if page and y + height > BODY_BOTTOM:
            pages.append(page)
            page, y = [], TOP_MARGIN
        for text, size, gap, _ in run:
            y += gap if page else 0.0
            page.append((y, text, size))
            y += size * LEADING
        index = last + 1
    pages.append(page)
    return pages


def _place_rail(items: list[tuple[str, float, float]]) -> list[tuple[float, str, float]]:
    placed: list[tuple[float, str, float]] = []
    y = TOP_MARGIN
    for text, size, gap in items:
        y += gap if placed else 0.0
        placed.append((y, text, size))
        y += size * LEADING
    if y > BODY_BOTTOM:
        raise ValueError(
            "the contact rail does not fit on page one; a LinkedIn export never "
            "continues it onto page two"
        )
    return placed


def build_linkedin_pdf(persona: dict[str, Any], path: Path) -> Path:
    rail = _place_rail(_rail_lines(persona))
    pages = _flow(_body_items(persona))

    path.parent.mkdir(parents=True, exist_ok=True)
    # invariant: no wall-clock creation date and no random document id, so the
    # same persona renders byte for byte the same export.
    canvas = pdfcanvas.Canvas(
        str(path), pagesize=(PAGE_W, PAGE_H), invariant=1, pageCompression=1
    )
    canvas.setTitle(f"Profile - {persona['identity']['name']}")
    canvas.setAuthor(persona["identity"]["name"])
    canvas.setSubject(persona["notice"])

    def draw(x: float, y_top: float, text: str, size: float) -> None:
        canvas.setFont(FONT, size)
        canvas.drawString(x, PAGE_H - y_top - size, text)

    for number, page in enumerate(pages, start=1):
        if number == 1:
            for y, text, size in rail:
                draw(RAIL_X, y, text, size)
        for y, text, size in page:
            draw(BODY_X, y, text, size)
        draw(FOOTER_X, FOOTER_Y, f"Page {number} of {len(pages)}", SIZE_FOOTER)
        canvas.showPage()
    canvas.save()
    return path


# ---------------------------------------------------------------------------
# Verification (FR-102, FR-103)
# ---------------------------------------------------------------------------


def verify(persona: dict[str, Any], cv_path: Path, pdf_path: Path) -> bool:
    """Read both documents back with the product's own parsers and merge them.

    A fixture is only worth having if the code under test can read it, and the
    planted conflicts are only worth planting if the merge raises them.  Both
    are cheap to check here and expensive to discover from a failing
    end-to-end run.
    """
    cv = parse_cv(cv_path)
    linkedin = parse_linkedin_pdf(pdf_path)
    merged = merge_documents(linkedin, cv)

    print()
    print("Parsed back (FR-102, FR-103):")
    photo = f"{len(cv.photo) / 1024:.1f} kB" if cv.photo else "MISSING"
    print(
        f"  linkedin_pdf : name={linkedin.sections['contact']['name']!r} "
        f"roles={len(linkedin.sections['experience'])} "
        f"top_skills={len(linkedin.sections['top_skills'])} "
        f"languages={len(linkedin.sections['languages'])} "
        f"education={len(linkedin.sections['education'])}"
    )
    print(
        f"  cv           : name={cv.sections['contact']['name']!r} "
        f"roles={len(cv.sections['experience'])} "
        f"skills={len(cv.sections['top_skills'])} "
        f"certifications={len(cv.sections['certifications'])} photo={photo}"
    )
    print(
        f"  merged       : roles={len(merged.sections['experience'])} "
        f"skills={len(merged.sections['top_skills'])} "
        f"conflicts={len(merged.conflicts)}"
    )
    for warning in linkedin.warnings + cv.warnings:
        print(f"  warning      : {warning}")

    ok = True
    print()
    print("Planted conflicts:")
    for planted in persona["planted_conflicts"]:
        match = next(
            (
                c for c in merged.conflicts
                if c["field_path"] == planted["expected_field_path"]
                and str(c["value_linkedin"]) == str(planted["value_linkedin"])
                and str(c["value_cv"]) == str(planted["value_cv"])
            ),
            None,
        )
        if match is not None:
            print(f"  found    {planted['kind']:<9} {planted['expected_field_path']}")
            continue
        ok = False
        loose = [
            c for c in merged.conflicts
            if str(c["value_cv"]) == str(planted["value_cv"])
        ]
        where = f" (raised instead at {loose[0]['field_path']})" if loose else ""
        print(
            f"  MISSING  {planted['kind']:<9} {planted['expected_field_path']} "
            f"linkedin={planted['value_linkedin']!r} cv={planted['value_cv']!r}{where}"
        )

    unexpected = [
        c for c in merged.conflicts
        if c["field_path"] not in {p["expected_field_path"] for p in persona["planted_conflicts"]}
    ]
    for conflict in unexpected:
        ok = False
        print(
            f"  UNPLANTED          {conflict['field_path']} "
            f"linkedin={conflict['value_linkedin']!r} cv={conflict['value_cv']!r}"
        )
    if not cv.photo:
        ok = False
        print("  MISSING  photo     the CV carries no extractable image")
    return ok


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def main() -> int:
    parser = argparse.ArgumentParser(description="Render the end-to-end persona documents")
    parser.add_argument("--persona", type=Path, default=PERSONA_PATH)
    parser.add_argument("--out", type=Path, default=OUT_DIR)
    parser.add_argument("--skip-verify", action="store_true")
    args = parser.parse_args()

    if not args.persona.exists():
        print(f"No persona at {args.persona}; run generate.py first.")
        return 2
    persona = json.loads(args.persona.read_text(encoding="utf-8"))

    avatar = build_avatar(persona, args.out / AVATAR_PATH.name)
    cv_path = build_cv(persona, args.out / CV_PATH.name, avatar)
    pdf_path = build_linkedin_pdf(persona, args.out / LINKEDIN_PATH.name)

    identity = persona["identity"]
    print(f"Persona    : {identity['name']} - {identity['headline']}")
    print(f"Avatar     : {avatar}  ({avatar.stat().st_size / 1024:.1f} kB)")
    print(
        f"CV         : {cv_path}  ({cv_path.stat().st_size / 1024:.1f} kB, "
        f"{len(persona['experience'])} roles, {len(persona['skills'])} skills, "
        f"{len(persona['skill_groups'])} skill families)"
    )
    print(
        f"LinkedIn   : {pdf_path}  ({pdf_path.stat().st_size / 1024:.1f} kB, "
        f"{len(persona['top_skills'])} top skills, {len(persona['languages'])} languages)"
    )

    if args.skip_verify:
        return 0
    return 0 if verify(persona, cv_path, pdf_path) else 1


if __name__ == "__main__":
    raise SystemExit(main())
